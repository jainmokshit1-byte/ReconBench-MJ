"""Synthetic end-to-end checks for upload, queue hand-off, and parsed document persistence."""

from __future__ import annotations

import hashlib
import io
import json
import os
import time
import zipfile
from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

import boto3
import pytest
import reconbench_app.jobs as job_module
import sqlalchemy as sa
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reconbench_app.api.application import create_app
from reconbench_app.authentication import (
    COOKIE_NAME,
    AuthContext,
    AuthenticationCheck,
)
from reconbench_app.configuration import canonical_hash
from reconbench_app.db.ids import new_id
from reconbench_app.db.metadata import metadata
from reconbench_app.db.seed_data import DEFAULT_CONFIG, REASON_CODES
from reconbench_app.export_service import ExportService
from reconbench_app.ingestion import (
    PARSER_VERSION,
    IngestionService,
    sniff_document_type,
    sniff_format,
)
from reconbench_app.ledger_service import LedgerRebuildService
from reconbench_app.mapping_service import MappingService
from reconbench_app.pdf_sandbox import page_image_hashes_isolated
from reconbench_app.settings import Settings
from reconbench_app.storage import LocalObjectStore, S3ObjectStore
from reconbench_app.vendor_resolution import VendorResolutionService
from reconbench_core.parsers.types import PageImageHash
from reconbench_core.schema import (
    DocLine,
    DocType,
    Document,
    ExtractionMethod,
    FieldEvidence,
    FieldOrigin,
    FieldSource,
    FieldValue,
    FieldVerification,
    FileFormat,
    InvoiceHeader,
    POHeader,
    ReasonCode,
)
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy.engine import Engine, make_url
from sqlalchemy.exc import IntegrityError

_CSV = (
    b"Invoice No,Invoice Date,Vendor Name,Vendor GSTIN,Description,HSN,UOM,Quantity,Rate,"
    b"Amount,Taxable Total,CGST,SGST,Grand Total\n"
    b"INV-2026-001,01/04/2026,Acme Components Pvt Ltd,27AAPFU0939F1ZV,Widget A,8471,NOS,"
    b"2,100,200,200,18,18,236\n"
)

_PO_CSV = (
    b"PO No,Order Date,Vendor Name,Vendor GSTIN,Currency,Taxable Total,Grand Total,Line No,"
    b"Description,HSN,UOM,Quantity,Rate,Amount\n"
    b"PO-2026-001,01/04/2026,Acme Components Pvt Ltd,27AAPFU0939F1ZV,INR,200,236,1,"
    b"Widget A,8471,NOS,2,100,200\n"
)

_LINKABLE_INVOICE_CSV = (
    b"Invoice No,Invoice Date,PO Ref,Vendor Name,Vendor GSTIN,Currency,Taxable Total,Grand Total,"
    b"Line No,Description,HSN,UOM,Quantity,Rate,Amount\n"
    b"INV-2026-002,02/04/2026,PO-2026-001,Acme Components Pvt Ltd,27AAPFU0939F1ZV,INR,"
    b"200,236,1,Widget A,8471,NOS,2,100,200\n"
)

_BETA_PO_CSV = (
    b"PO No,Order Date,Vendor Name,Vendor GSTIN,Currency,Taxable Total,Grand Total,"
    b"PO Status,Line No,"
    b"Description,HSN,UOM,Quantity,Rate,Amount\n"
    b"PO-2026-777,03/04/2026,Acme Components Pvt Ltd,27AAPFU0939F1ZV,INR,100,118,open,1,"
    b"Widget B,8471,NOS,2,50,100\n"
)

_BETA_INVOICE_CSV = (
    b"Invoice No,Invoice Date,PO Ref,Vendor Name,Vendor GSTIN,Currency,Taxable Total,Grand Total,"
    b"CGST,SGST,Line No,Description,HSN,UOM,Quantity,Rate,Amount\n"
    b"APR-26-INV-901,03/04/2026,PO-2026-777,Acme Components Pvt Ltd,27AAPFU0939F1ZV,INR,"
    b"100,118,9,9,1,Widget B,8471,NOS,2,50,100\n"
)

_CONTENT_PO_CSV = (
    _BETA_PO_CSV.replace(b"PO-2026-777", b"PO-2026-778")
    .replace(b"Widget B", b"Widget C Compact")
    .replace(b",2,50,100\n", b",2,70,140\n")
    .replace(b",100,118,open", b",140,165.2,open")
)

_CONTENT_INVOICE_NO_PO_REF_CSV = (
    _BETA_INVOICE_CSV.replace(b"APR-26-INV-901", b"APR-26-INV-903")
    .replace(b"03/04/2026,PO-2026-777", b"06/04/2026,")
    .replace(b"Widget B", b"Widget C Compact")
    .replace(b",100,118,9,9,1,", b",140,165.2,12.6,12.6,1,")
    .replace(b",2,50,100\n", b",2,70,140\n")
)

_BETA_GRN_CSV = (
    b"GRN No,GRN Date,PO Ref,Vendor Name,Vendor GSTIN,Currency,Line No,Description,HSN,UOM,"
    b"Received Qty,Accepted Qty,Rejected Qty,Rate,Amount\n"
    b"GRN-2026-777,05/04/2026,PO-2026-777,Acme Components Pvt Ltd,27AAPFU0939F1ZV,INR,1,"
    b"Widget B,8471,NOS,2,2,0,50,100\n"
)

_CONTENT_GRN_CSV = (
    _BETA_GRN_CSV.replace(b"GRN-2026-777", b"GRN-2026-778")
    .replace(b"PO-2026-777", b"PO-2026-778")
    .replace(b"Widget B", b"Widget C Compact")
    .replace(b",2,2,0,50,100\n", b",2,2,0,70,140\n")
)


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    result = io.BytesIO()
    with zipfile.ZipFile(result, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    return result.getvalue()


def _pdf_bytes(*lines: str) -> bytes:
    stream = io.BytesIO()
    page = canvas.Canvas(stream, pagesize=letter)
    page.setFont("Helvetica", 10)
    for index, line in enumerate(lines):
        page.drawString(36, 748 - index * 18, line)
    page.save()
    return stream.getvalue()


def _scanned_pdf_bytes() -> bytes:
    stream = io.BytesIO()
    page = canvas.Canvas(stream, pagesize=letter)
    page.showPage()
    page.save()
    return stream.getvalue()


class _FakeQueue:
    def __init__(self) -> None:
        self.jobs: list[tuple[str, UUID]] = []

    def enqueue(self, task_name: str, *, file_id: UUID, connection: sa.Connection) -> int:
        assert connection.in_transaction()
        self.jobs.append((task_name, file_id))
        return len(self.jobs)

    def enqueue_reconciliation(self, *, run_id: UUID, connection: sa.Connection) -> int:
        assert connection.in_transaction()
        self.jobs.append(("evaluate_cases", run_id))
        return len(self.jobs)

    def enqueue_pending_recomputes(self, *, workspace_id: UUID, connection: sa.Connection) -> int:
        del workspace_id
        assert connection.in_transaction()
        return 0

    def enqueue_export(self, *, export_id: UUID, connection: sa.Connection) -> int:
        assert connection.in_transaction()
        self.jobs.append(("build_export", export_id))
        return len(self.jobs)


class _FakeObjectStoreClient:
    def list_buckets(self) -> dict[str, list[object]]:
        return {"Buckets": []}


@pytest.fixture(scope="module")
def ingestion_engine() -> Iterator[Engine]:
    dsn = os.environ.get("RB_PHASE4_TEST_DATABASE_URL", "").strip()
    variable = "RB_PHASE4_TEST_DATABASE_URL"
    if not dsn:
        dsn = os.environ.get("RB_PHASE2_TEST_DATABASE_URL", "").strip()
        variable = "RB_PHASE2_TEST_DATABASE_URL"
    if not dsn:
        pytest.skip("RB_PHASE4_TEST_DATABASE_URL or RB_PHASE2_TEST_DATABASE_URL is required")
    database_name = make_url(dsn).database or ""
    expected_prefix = (
        "reconbench_phase4" if variable == "RB_PHASE4_TEST_DATABASE_URL" else "reconbench_phase2_"
    )
    if not database_name.startswith(expected_prefix):
        raise RuntimeError("ingestion tests require a disposable ReconBench Phase 2/4 database")
    engine = sa.create_engine(dsn, pool_size=8, max_overflow=8)
    try:
        with engine.connect() as connection:
            server_version = int(connection.scalar(sa.text("SHOW server_version_num")))
            if server_version // 10_000 != 16:
                raise RuntimeError("ingestion integration tests require PostgreSQL 16")
        yield engine
    finally:
        engine.dispose()


def _seed_workspace(engine: Engine) -> tuple[UUID, UUID]:
    workspace = metadata.tables["workspace"]
    app_user = metadata.tables["app_user"]
    config_version = metadata.tables["config_version"]
    config_active = metadata.tables["config_active"]
    reason_code = metadata.tables["reason_code"]
    workspace_id = new_id()
    user_id = new_id()
    config_id = new_id()
    with engine.begin() as connection:
        connection.execute(
            sa.insert(workspace).values(id=workspace_id, name="Synthetic Ingestion Workspace")
        )
        connection.execute(
            sa.insert(app_user).values(
                id=user_id,
                workspace_id=workspace_id,
                email=f"ingestion-{workspace_id}@example.invalid",
                display_name="Synthetic Ingestion Manager",
                role="manager",
                password_hash="synthetic-unusable-hash",
            )
        )
        connection.execute(
            sa.insert(config_version).values(
                id=config_id,
                workspace_id=workspace_id,
                seq=1,
                body=DEFAULT_CONFIG,
                engine_hash=canonical_hash(DEFAULT_CONFIG["engine"]),
                body_hash=canonical_hash(DEFAULT_CONFIG),
                change_note="Synthetic ingestion baseline",
                created_by=user_id,
            )
        )
        connection.execute(
            sa.insert(config_active).values(
                id=new_id(),
                workspace_id=workspace_id,
                config_version_id=config_id,
                activated_by=user_id,
                version=1,
            )
        )
        for item in REASON_CODES:
            connection.execute(
                sa.insert(reason_code).values(id=new_id(), workspace_id=workspace_id, **item)
            )
    return workspace_id, user_id


def test_csv_upload_is_queued_parsed_and_duplicate_upload_is_idempotent(
    ingestion_engine: Engine,
    tmp_path: Any,
    monkeypatch: pytest.MonkeyPatch,
    request: pytest.FixtureRequest,
) -> None:
    workspace_id, user_id = _seed_workspace(ingestion_engine)

    s3_endpoint = os.environ.get("RB_PHASE4_TEST_S3_ENDPOINT_URL", "").strip()
    s3_bucket = "synthetic-ingestion"
    test_s3_client: Any = _FakeObjectStoreClient()
    if s3_endpoint:
        s3_client = boto3.client(
            "s3",
            endpoint_url=s3_endpoint,
            region_name=os.environ.get("RB_PHASE4_TEST_S3_REGION", "us-east-1"),
            aws_access_key_id=os.environ["RB_PHASE4_TEST_S3_ACCESS_KEY"],
            aws_secret_access_key=os.environ["RB_PHASE4_TEST_S3_SECRET_KEY"],
        )
        s3_bucket = f"rb-phase4-{workspace_id.hex}"
        s3_client.create_bucket(Bucket=s3_bucket)

        def cleanup_s3_bucket() -> None:
            pages = s3_client.get_paginator("list_objects_v2").paginate(Bucket=s3_bucket)
            for result in pages:
                objects = [{"Key": item["Key"]} for item in result.get("Contents", [])]
                if objects:
                    s3_client.delete_objects(
                        Bucket=s3_bucket,
                        Delete={"Objects": objects, "Quiet": True},
                    )
            s3_client.delete_bucket(Bucket=s3_bucket)

        request.addfinalizer(cleanup_s3_bucket)
        monkeypatch.setattr(boto3, "client", lambda *args, **kwargs: s3_client)
        storage = S3ObjectStore(s3_client, s3_bucket)
        test_s3_client = s3_client
    else:
        monkeypatch.setattr(boto3, "client", lambda *args, **kwargs: test_s3_client)
        storage = LocalObjectStore(tmp_path / "objects")

    settings = Settings(
        _env_file=None,
        environment="ci",
        database_url=ingestion_engine.url.render_as_string(hide_password=False),
        s3_endpoint_url=s3_endpoint or "http://127.0.0.1:1",
        s3_region=os.environ.get("RB_PHASE4_TEST_S3_REGION", "us-east-1"),
        s3_bucket=s3_bucket,
        s3_api_access_key="synthetic-access",
        s3_api_secret_key="synthetic-secret",
        public_base_url="https://example.invalid",
        argon2_time_cost=1,
        argon2_memory_cost_kib=8192,
        argon2_parallelism=1,
    )
    app = create_app(settings)
    queue = _FakeQueue()
    service = IngestionService(ingestion_engine, storage, queue)
    app.state.ingestion_service = service
    context = AuthContext(
        session_id=new_id(),
        workspace_id=workspace_id,
        user_id=user_id,
        display_name="Synthetic Ingestion Manager",
        role="manager",
        csrf_token="synthetic-csrf",
    )
    app.state.auth_service.authenticate = lambda *args, **kwargs: AuthenticationCheck(context, True)

    with TestClient(app, base_url="https://testserver") as client:
        app.state.file_storage = storage
        client.cookies.set(COOKIE_NAME, "synthetic-session", path="/api/v1")
        # G-12 tracker: an empty workspace starts at Ingest.
        empty_workflow = client.get("/api/v1/workflow-status?from=2026-04-01&to=2026-04-30")
        assert empty_workflow.status_code == 200, empty_workflow.text
        assert empty_workflow.json()["next_step"] == "ingest"
        assert not any(step["complete"] for step in empty_workflow.json()["steps"])
        batch_response = client.post(
            "/api/v1/batches",
            json={"label": "Synthetic April 2026"},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert batch_response.status_code == 201, batch_response.text
        batch_id = UUID(batch_response.json()["id"])

        upload = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=_CSV,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "invoice.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert upload.status_code == 201, upload.text
        file_id = UUID(upload.json()["file_id"])
        assert upload.json()["format"] == "csv"
        assert upload.json()["status"] == "staged"

        process = client.post(
            f"/api/v1/batches/{batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        assert queue.jobs == [("ingest_file", file_id)]
        service.ingest_file(file_id, job_id=1)
        assert queue.jobs[-1] == ("parse_file", file_id)
        service.ingest_file(file_id, job_id=1)
        assert queue.jobs.count(("parse_file", file_id)) == 1
        object_get = storage.get

        def simulate_worker_exit(*args: Any, **kwargs: Any) -> Any:
            del args, kwargs
            raise KeyboardInterrupt

        monkeypatch.setattr(storage, "get", simulate_worker_exit)
        with pytest.raises(KeyboardInterrupt):
            service.parse_file(file_id, job_id=2)
        monkeypatch.setattr(storage, "get", object_get)
        service.parse_file(file_id, job_id=2)
        # A committed parse whose queue acknowledgement is lost is an idempotent no-op.
        service.parse_file(file_id, job_id=2)
        assert queue.jobs.count(("resolve_vendors", file_id)) == 1

        file_response = client.get(f"/api/v1/files/{file_id}")
        assert file_response.status_code == 200
        assert file_response.json()["status"] == "parsed"
        assert file_response.json()["docs_new"] == 1

        with ingestion_engine.begin() as connection:
            connection.execute(
                sa.update(metadata.tables["source_file"])
                .where(metadata.tables["source_file"].c.id == file_id)
                .values(parser_version="phase4-parser-2")
            )
        outdated_file = client.get(f"/api/v1/files/{file_id}")
        assert outdated_file.json()["retry_available"] is True
        assert outdated_file.json()["retry_reason"] == "PARSER_VERSION_OUTDATED"
        reprocess = client.post(
            f"/api/v1/files/{file_id}/retry",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert reprocess.status_code == 202, reprocess.text
        assert queue.jobs[-1] == ("ingest_file", file_id)
        service.ingest_file(file_id, job_id=3)
        service.parse_file(file_id, job_id=4)
        refreshed_file = client.get(f"/api/v1/files/{file_id}")
        assert refreshed_file.json()["status"] == "parsed"
        assert refreshed_file.json()["parser_version"] == PARSER_VERSION
        assert refreshed_file.json()["retry_available"] is False

        documents = client.get(f"/api/v1/documents?batch_id={batch_id}")
        assert documents.status_code == 200
        assert len(documents.json()) == 1
        source_documents = client.get(
            f"/api/v1/documents?batch_id={batch_id}&source_file_id={file_id}"
        )
        assert source_documents.status_code == 200
        assert [row["id"] for row in source_documents.json()] == [documents.json()[0]["id"]]
        document_id = UUID(documents.json()[0]["id"])
        details = client.get(f"/api/v1/documents/{document_id}")
        assert details.status_code == 200
        assert len(details.json()["lines"]) == 1
        assert any(
            evidence["field_name"] == "vendor_name" for evidence in details.json()["field_evidence"]
        )
        assert details.json()["document"]["id"] == str(document_id)
        assert details.json()["version"]["document_id"] == str(document_id)
        assert details.json()["document"]["doc_no"] == "INV-2026-001"
        assert details.json()["version"]["doc_date"] == "2026-04-01"
        assert any(
            event["stage"] == "upserted" and event["document_id"] == str(document_id)
            for event in details.json()["processing_log"]
        )

        vendor_resolution = VendorResolutionService(ingestion_engine)
        with ThreadPoolExecutor(max_workers=2) as executor:
            concurrent_resolutions = list(
                executor.map(lambda _: vendor_resolution.resolve_workspace(workspace_id), range(2))
            )
        assert sum(result.documents_seen for result in concurrent_resolutions) == 2
        assert sum(result.documents_assigned for result in concurrent_resolutions) == 2
        assert sum(result.vendors_created for result in concurrent_resolutions) == 1
        assert sum(result.registrations_added for result in concurrent_resolutions) == 1
        with ingestion_engine.connect() as connection:
            resolved_vendor_id = connection.scalar(
                sa.select(metadata.tables["document"].c.vendor_id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.id == document_id,
                )
            )
            assert resolved_vendor_id is not None
            assert (
                connection.scalar(
                    sa.select(metadata.tables["vendor"].c.anchor).where(
                        metadata.tables["vendor"].c.workspace_id == workspace_id,
                        metadata.tables["vendor"].c.id == resolved_vendor_id,
                    )
                )
                == "gstin"
            )
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["vendor_gstin"])
                    .where(
                        metadata.tables["vendor_gstin"].c.workspace_id == workspace_id,
                        metadata.tables["vendor_gstin"].c.vendor_id == resolved_vendor_id,
                        metadata.tables["vendor_gstin"].c.gstin == "27AAPFU0939F1ZV",
                    )
                )
                == 1
            )
            vendor_input_seq = connection.scalar(
                sa.select(metadata.tables["vendor"].c.input_seq).where(
                    metadata.tables["vendor"].c.workspace_id == workspace_id,
                    metadata.tables["vendor"].c.id == resolved_vendor_id,
                )
            )
        second_resolution = vendor_resolution.resolve_workspace(workspace_id)
        assert second_resolution.vendors_created == 0
        assert second_resolution.registrations_added == 0
        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(metadata.tables["vendor"].c.input_seq).where(
                        metadata.tables["vendor"].c.workspace_id == workspace_id,
                        metadata.tables["vendor"].c.id == resolved_vendor_id,
                    )
                )
                == vendor_input_seq
            )

        def upload_and_parse(
            zone: str,
            name: str,
            content: bytes,
            job_base: int,
            mapping: dict[str, str] | None = None,
        ) -> UUID:
            response = client.post(
                f"/api/v1/files?batch_id={batch_id}&zone={zone}",
                content=content,
                headers={
                    "Content-Type": "application/octet-stream",
                    "X-File-Name": name,
                    "X-CSRF-Token": "synthetic-csrf",
                },
            )
            assert response.status_code == 201, response.text
            uploaded_id = UUID(response.json()["file_id"])
            processed = client.post(
                f"/api/v1/batches/{batch_id}/process",
                headers={"X-CSRF-Token": "synthetic-csrf"},
            )
            assert processed.status_code == 202, processed.text
            service.ingest_file(uploaded_id, job_id=job_base)
            service.parse_file(uploaded_id, job_id=job_base + 1)
            if mapping is not None:
                file_state = client.get(f"/api/v1/files/{uploaded_id}")
                assert file_state.status_code == 200
                assert file_state.json()["error_code"] == "CSV_MAPPING_REQUIRED"
                saved = client.post(
                    f"/api/v1/files/{uploaded_id}/column-mapping",
                    json={"column_mapping": mapping},
                    headers={"X-CSRF-Token": "synthetic-csrf"},
                )
                assert saved.status_code == 202, saved.text
                service.ingest_file(uploaded_id, job_id=job_base + 2)
                service.parse_file(uploaded_id, job_id=job_base + 3)
            return uploaded_id

        upload_and_parse(
            "po", "purchase-order.csv", _PO_CSV.replace(b"Widget A", b"Widget A MOD-100"), 30
        )
        upload_and_parse(
            "invoice",
            "invoice-with-po.csv",
            _LINKABLE_INVOICE_CSV.replace(b"Widget A", b"Widget A MOD-100"),
            32,
            mapping={
                "PO Ref": "po_refs",
                "Invoice No": "doc_no",
                "Invoice Date": "doc_date",
            },
        )
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service = MappingService(ingestion_engine)
        mapping_summary = mapping_service.resolve_workspace(workspace_id)
        assert mapping_summary.linked >= 1
        assert mapping_summary.link_versions_written >= 1
        with ingestion_engine.connect() as connection:
            po_document_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_type == "po",
                    metadata.tables["document"].c.doc_no == "PO-2026-001",
                )
            )
        assert po_document_id is not None
        ledger_service = LedgerRebuildService(ingestion_engine)
        ledger_first = ledger_service.rebuild_po(workspace_id, po_document_id)
        assert ledger_first == {"lines": 1, "entries": 2, "changed": True}
        ledger_replay = ledger_service.rebuild_po(workspace_id, po_document_id)
        assert ledger_replay == {"lines": 1, "entries": 0, "changed": False}
        with ThreadPoolExecutor(max_workers=2) as executor:
            concurrent_rebuilds = list(
                executor.map(
                    lambda _: ledger_service.rebuild_po(workspace_id, po_document_id),
                    range(2),
                )
            )
        assert concurrent_rebuilds == [
            {"lines": 1, "entries": 0, "changed": False},
            {"lines": 1, "entries": 0, "changed": False},
        ]
        with ingestion_engine.connect() as connection:
            ledger_row = connection.execute(
                sa.select(
                    metadata.tables["po_line_ledger"].c.ordered,
                    metadata.tables["po_line_ledger"].c.billed,
                    metadata.tables["po_line_ledger"].c.received_accepted,
                    metadata.tables["po_line_ledger"].c.open_qty,
                ).where(
                    metadata.tables["po_line_ledger"].c.workspace_id == workspace_id,
                    metadata.tables["po_line_ledger"].c.po_document_id == po_document_id,
                )
            ).one()
            assert ledger_row == (
                Decimal("2.0000"),
                Decimal("4.0000"),
                Decimal("0.0000"),
                Decimal("-2.0000"),
            )

        dirty_scope = metadata.tables["dirty_scope"]
        with ingestion_engine.begin() as connection:
            pending_po_scope = connection.scalar(
                sa.select(dirty_scope.c.id).where(
                    dirty_scope.c.workspace_id == workspace_id,
                    dirty_scope.c.scope_kind == "po",
                    dirty_scope.c.scope_id == po_document_id,
                    dirty_scope.c.processed_at.is_(None),
                )
            )
            if pending_po_scope is None:
                pending_po_scope = new_id()
                connection.execute(
                    sa.insert(dirty_scope).values(
                        id=pending_po_scope,
                        workspace_id=workspace_id,
                        scope_kind="po",
                        scope_id=po_document_id,
                        reason="synthetic_worker_interruption",
                    )
                )
        po_line_ledger = metadata.tables["po_line_ledger"]
        ledger_entry = metadata.tables["ledger_entry"]
        with ingestion_engine.connect() as connection:
            before_ledger_row = connection.execute(
                sa.select(po_line_ledger.c.input_hash, po_line_ledger.c.billed).where(
                    po_line_ledger.c.workspace_id == workspace_id,
                    po_line_ledger.c.po_document_id == po_document_id,
                )
            ).one()
            before_entry_count = connection.scalar(
                sa.select(sa.func.count())
                .select_from(
                    ledger_entry.join(
                        po_line_ledger,
                        ledger_entry.c.po_line_ledger_id == po_line_ledger.c.id,
                    )
                )
                .where(
                    po_line_ledger.c.workspace_id == workspace_id,
                    po_line_ledger.c.po_document_id == po_document_id,
                )
            )
        original_rebuild = ledger_service.rebuild_po_connection

        def interrupt_after_rebuild(
            connection: sa.Connection, current_workspace_id: UUID, current_po_id: UUID
        ) -> dict[str, int | bool]:
            result = original_rebuild(connection, current_workspace_id, current_po_id)
            if current_po_id == po_document_id:
                raise RuntimeError("synthetic worker termination")
            return result

        monkeypatch.setattr(ledger_service, "rebuild_po_connection", interrupt_after_rebuild)
        with pytest.raises(RuntimeError, match="synthetic worker termination"):
            ledger_service.drain_dirty_scopes(workspace_id)
        monkeypatch.setattr(ledger_service, "rebuild_po_connection", original_rebuild)
        with ingestion_engine.connect() as connection:
            pending_after_interruption = connection.scalar(
                sa.select(dirty_scope.c.processed_at).where(dirty_scope.c.id == pending_po_scope)
            )
            after_ledger_row = connection.execute(
                sa.select(po_line_ledger.c.input_hash, po_line_ledger.c.billed).where(
                    po_line_ledger.c.workspace_id == workspace_id,
                    po_line_ledger.c.po_document_id == po_document_id,
                )
            ).one()
            after_entry_count = connection.scalar(
                sa.select(sa.func.count())
                .select_from(
                    ledger_entry.join(
                        po_line_ledger,
                        ledger_entry.c.po_line_ledger_id == po_line_ledger.c.id,
                    )
                )
                .where(
                    po_line_ledger.c.workspace_id == workspace_id,
                    po_line_ledger.c.po_document_id == po_document_id,
                )
            )
        assert pending_after_interruption is None
        assert after_ledger_row == before_ledger_row
        assert after_entry_count == before_entry_count
        drained = ledger_service.drain_dirty_scopes(workspace_id)
        assert drained["scopes"] >= 1
        assert drained["pos"] == 1
        assert drained["changed"] == 0
        drained_again = ledger_service.drain_dirty_scopes(workspace_id)
        assert drained_again == {"scopes": 0, "pos": 0, "changed": 0, "entries": 0}
        with ingestion_engine.connect() as connection:
            linked_invoice_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_no == "INV-2026-002",
                )
            )
            assert linked_invoice_id is not None
            link_row = connection.execute(
                sa.select(
                    metadata.tables["link"].c.method,
                    metadata.tables["link"].c.state,
                    metadata.tables["link"].c.confidence,
                )
                .select_from(
                    metadata.tables["link_current"].join(
                        metadata.tables["link"],
                        metadata.tables["link_current"].c.current_link_id
                        == metadata.tables["link"].c.id,
                    )
                )
                .where(
                    metadata.tables["link_current"].c.workspace_id == workspace_id,
                    metadata.tables["link_current"].c.from_document_id == linked_invoice_id,
                )
            ).one()
            assert link_row.method == "labeled"
            assert link_row.state == "auto"
            assert link_row.confidence == Decimal("0.99")
            link_count_before_replay = connection.scalar(
                sa.select(sa.func.count())
                .select_from(metadata.tables["link"])
                .where(
                    metadata.tables["link"].c.workspace_id == workspace_id,
                    metadata.tables["link"].c.from_document_id == linked_invoice_id,
                )
            )
        replay_summary = mapping_service.resolve_workspace(workspace_id)
        assert replay_summary.link_versions_written == 0
        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["link"])
                    .where(
                        metadata.tables["link"].c.workspace_id == workspace_id,
                        metadata.tables["link"].c.from_document_id == linked_invoice_id,
                    )
                )
                == link_count_before_replay
            )

        upload_and_parse(
            "po",
            "beta-purchase-order.csv",
            _BETA_PO_CSV.replace(b"Widget B", b"Widget B MOD-777"),
            80,
        )
        upload_and_parse(
            "invoice",
            "beta-invoice.csv",
            _BETA_INVOICE_CSV.replace(b"Widget B", b"Widget B MOD-777"),
            82,
            mapping={
                "PO Ref": "po_refs",
                "Invoice No": "doc_no",
                "Invoice Date": "doc_date",
            },
        )
        vendor_resolution.resolve_workspace(workspace_id)
        beta_mapping = mapping_service.resolve_workspace(workspace_id)
        assert beta_mapping.linked >= 1
        with ingestion_engine.connect() as connection:
            beta_invoice_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_no == "APR-26-INV-901",
                )
            )
            beta_po_document_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_type == "po",
                    metadata.tables["document"].c.doc_no == "PO-2026-777",
                )
            )
        assert beta_invoice_id is not None
        assert beta_po_document_id is not None

        app.state.job_queue = queue
        run_response = client.post(
            "/api/v1/runs",
            json={"run_date": "2026-04-30"},
            headers={
                "Idempotency-Key": "replay-test-run-create",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert run_response.status_code == 202, run_response.text
        initial_run_id = UUID(run_response.json()["id"])
        app.state.reconciliation_service.evaluate_run(initial_run_id)
        case_current = metadata.tables["case_current"]
        with ingestion_engine.connect() as connection:
            initial_invoice_states = {
                row.invoice_id: (row.effective_status, row.run_id)
                for row in connection.execute(
                    sa.select(
                        case_current.c.invoice_id,
                        case_current.c.effective_status,
                        case_current.c.run_id,
                    ).where(case_current.c.workspace_id == workspace_id)
                )
            }
            beta_status, beta_initial_run = initial_invoice_states[beta_invoice_id]
        assert beta_status == "pending"
        assert beta_initial_run == initial_run_id

        grn_file_id = upload_and_parse(
            "grn",
            "beta-receipt.csv",
            _BETA_GRN_CSV.replace(b"Widget B", b"Widget B MOD-777"),
            90,
            mapping={"PO Ref": "po_refs", "GRN No": "doc_no", "GRN Date": "doc_date"},
        )
        assert grn_file_id
        vendor_resolution.resolve_workspace(workspace_id)
        grn_mapping = mapping_service.resolve_workspace(workspace_id)
        assert grn_mapping.linked >= 1
        with ingestion_engine.connect() as connection:
            grn_document_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_type == "grn",
                    metadata.tables["document"].c.doc_no == "GRN-2026-777",
                )
            )
        assert grn_document_id is not None
        with ThreadPoolExecutor(max_workers=2) as executor:
            concurrent_beta_rebuilds = list(
                executor.map(
                    lambda _: ledger_service.rebuild_po(workspace_id, beta_po_document_id),
                    range(2),
                )
            )
        assert sorted(bool(result["changed"]) for result in concurrent_beta_rebuilds) == [
            False,
            True,
        ]
        assert all(result["lines"] == 1 for result in concurrent_beta_rebuilds)
        monkeypatch.setattr(
            job_module,
            "ProcrastinateJobQueue",
            lambda _tasks, **_kwargs: queue,
        )
        worker_app, _worker_queue = job_module.register_ingestion_tasks(
            settings, ingestion_engine, test_s3_client
        )
        queued_before_grn_recheck = len(queue.jobs)
        evaluation_jobs_before = [job for job in queue.jobs if job[0] == "evaluate_cases"]
        worker_app.tasks["reconbench.link_vendor"].func(None, str(grn_file_id))
        auto_recheck_jobs = queue.jobs[queued_before_grn_recheck:]
        auto_runs = [
            run_id for task_name, run_id in auto_recheck_jobs if task_name == "evaluate_cases"
        ]
        assert len(auto_runs) == 1
        worker_app.tasks["reconbench.link_vendor"].func(None, str(grn_file_id))
        assert [job for job in queue.jobs if job[0] == "evaluate_cases"] == (
            evaluation_jobs_before + [("evaluate_cases", auto_runs[0])]
        )
        worker_app.close()
        app.state.reconciliation_service.evaluate_run(auto_runs[0])
        with ingestion_engine.connect() as connection:
            beta_after_grn = connection.execute(
                sa.select(case_current.c.effective_status, case_current.c.run_id).where(
                    case_current.c.workspace_id == workspace_id,
                    case_current.c.invoice_id == beta_invoice_id,
                )
            ).one()
            other_case_after_grn = connection.execute(
                sa.select(case_current.c.effective_status, case_current.c.run_id).where(
                    case_current.c.workspace_id == workspace_id,
                    case_current.c.invoice_id == linked_invoice_id,
                )
            ).one()
            auto_run = connection.execute(
                sa.select(
                    metadata.tables["recon_run"].c.trigger_kind,
                    metadata.tables["recon_run"].c.scope,
                ).where(metadata.tables["recon_run"].c.id == auto_runs[0])
            ).one()
        assert beta_after_grn.run_id == auto_runs[0]
        assert beta_after_grn.effective_status in {"matched", "matched_flags"}
        assert other_case_after_grn.run_id == initial_run_id
        assert auto_run.trigger_kind == "auto_pending"
        assert auto_run.scope["invoice_ids"] == [str(beta_invoice_id)]

        kpi_response = client.get("/api/v1/kpis?from=2026-04-01&to=2026-04-30")
        assert kpi_response.status_code == 200, kpi_response.text
        with ingestion_engine.connect() as connection:
            independent_kpis = (
                connection.execute(
                    sa.text(
                        """
                    SELECT count(*) AS invoice_count,
                           round(100.0 * count(*) FILTER (
                               WHERE effective_status IN ('matched', 'matched_flags')
                           ) / NULLIF(count(*), 0), 2) AS matched_pct,
                           round(100.0 * count(*) FILTER (
                               WHERE computed_status = 'matched' AND NOT human_touched
                           ) / NULLIF(count(*), 0), 2) AS straight_through_pct,
                           count(*) FILTER (WHERE effective_status = 'pending') AS pending_count,
                           count(*) FILTER (
                               WHERE effective_status = 'exception'
                           ) AS exception_count,
                           count(*) FILTER (WHERE effective_status = 'unlinked') AS unlinked_count,
                           COALESCE(sum(at_risk), 0) AS at_risk_total,
                           COALESCE(sum(at_risk_gst), 0) AS at_risk_gst_total,
                           COALESCE(sum(blocked_value), 0) AS blocked_total,
                           (SELECT COALESCE(sum(grni_value), 0)
                              FROM v_grni_now
                             WHERE workspace_id = :workspace_id) AS grni_total
                      FROM v_case_current
                     WHERE workspace_id = :workspace_id
                       AND (invoice_date IS NULL
                            OR invoice_date BETWEEN DATE '2026-04-01' AND DATE '2026-04-30')
                    """
                    ),
                    {"workspace_id": workspace_id},
                )
                .mappings()
                .one()
            )
        api_kpis = kpi_response.json()
        for metric in (
            "invoice_count",
            "matched_pct",
            "straight_through_pct",
            "pending_count",
            "exception_count",
            "unlinked_count",
            "at_risk_total",
            "at_risk_gst_total",
            "blocked_total",
            "grni_total",
        ):
            if metric.endswith("_count"):
                assert api_kpis[metric] == independent_kpis[metric]
            else:
                assert Decimal(str(api_kpis[metric])) == Decimal(str(independent_kpis[metric]))

        # Hold a parse transaction immediately after its document-version write.
        # A run started while that row is uncommitted must see either the old
        # snapshot or the fully committed invoice, never a partial document.
        race_content = _CSV.replace(b"INV-2026-001", b"INV-2026-RACE")
        race_upload = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=race_content,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "invoice-ingest-run-race.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert race_upload.status_code == 201, race_upload.text
        race_file_id = UUID(race_upload.json()["file_id"])
        race_process = client.post(
            f"/api/v1/batches/{batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert race_process.status_code == 202, race_process.text
        service.ingest_file(race_file_id, job_id=132)

        advisory_key = (int(workspace_id.hex[:8], 16) & 0x7FFFFFFF) or 1
        trigger_suffix = workspace_id.hex
        trigger_function = f"rb_phase4_pause_parse_{trigger_suffix}"
        trigger_name = f"rb_phase4_pause_version_{trigger_suffix}"
        with ingestion_engine.begin() as connection:
            connection.exec_driver_sql(
                f"""
                CREATE FUNCTION {trigger_function}() RETURNS trigger LANGUAGE plpgsql AS $$
                BEGIN
                    PERFORM pg_advisory_xact_lock({advisory_key});
                    PERFORM pg_sleep(2);
                    RETURN NEW;
                END;
                $$
                """
            )
            connection.exec_driver_sql(
                f"""
                CREATE TRIGGER {trigger_name} BEFORE INSERT ON document_version
                FOR EACH ROW EXECUTE FUNCTION {trigger_function}()
                """
            )
        try:
            with ThreadPoolExecutor(max_workers=1) as executor:
                parse_future = executor.submit(lambda: service.parse_file(race_file_id, job_id=133))
                lock_seen = False
                lock_deadline = time.monotonic() + 8
                while time.monotonic() < lock_deadline:
                    if parse_future.done():
                        parse_future.result()
                        break
                    with ingestion_engine.connect() as connection:
                        lock_available = connection.scalar(
                            sa.text("SELECT pg_try_advisory_xact_lock(:key)"),
                            {"key": advisory_key},
                        )
                    if lock_available is False:
                        lock_seen = True
                        break
                    time.sleep(0.05)
                assert lock_seen, "parse transaction did not reach the version-write barrier"
                race_run = client.post(
                    "/api/v1/runs",
                    json={"run_date": "2026-04-30"},
                    headers={
                        "Idempotency-Key": "phase4-ingest-run-race",
                        "X-CSRF-Token": "synthetic-csrf",
                    },
                )
                assert race_run.status_code == 202, race_run.text
                race_run_id = UUID(race_run.json()["id"])
                with ingestion_engine.connect() as connection:
                    assert (
                        connection.scalar(
                            sa.select(
                                sa.exists().where(
                                    metadata.tables["document"].c.workspace_id == workspace_id,
                                    metadata.tables["document"].c.doc_no == "INV-2026-RACE",
                                )
                            )
                        )
                        is False
                    )
                parse_future.result(timeout=15)
        finally:
            with ingestion_engine.begin() as connection:
                connection.exec_driver_sql(
                    f"DROP TRIGGER IF EXISTS {trigger_name} ON document_version"
                )
                connection.exec_driver_sql(f"DROP FUNCTION IF EXISTS {trigger_function}()")

        app.state.reconciliation_service.evaluate_run(race_run_id)
        race_run_state = client.get(f"/api/v1/runs/{race_run_id}")
        assert race_run_state.status_code == 200, race_run_state.text
        assert race_run_state.json()["status"] == "succeeded"
        assert race_run_state.json()["done_cases"] == race_run_state.json()["total_cases"]
        with ingestion_engine.connect() as connection:
            race_invoice_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_no == "INV-2026-RACE",
                )
            )
            assert race_invoice_id is not None
            race_case_count = connection.scalar(
                sa.select(sa.func.count())
                .select_from(metadata.tables["recon_case"])
                .where(
                    metadata.tables["recon_case"].c.workspace_id == workspace_id,
                    metadata.tables["recon_case"].c.run_id == race_run_id,
                    metadata.tables["recon_case"].c.invoice_id == race_invoice_id,
                )
            )
        assert race_case_count == 1

        # A fresh all-invoice run after the ingest race becomes the common
        # anchor for Home KPIs and the export workbook.
        full_run_response = client.post(
            "/api/v1/runs",
            json={"run_date": "2026-04-30"},
            headers={
                "Idempotency-Key": "phase4-full-export-run",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert full_run_response.status_code == 202, full_run_response.text
        full_run_id = UUID(full_run_response.json()["id"])
        app.state.reconciliation_service.evaluate_run(full_run_id)

        # Manual Recon evaluates an operator-selected invoice/PO/GRN set without
        # creating a persistent reconciliation run or mutating the automatic case.
        with ingestion_engine.connect() as connection:
            beta_case_id, beta_case_version = connection.execute(
                sa.select(case_current.c.case_id, case_current.c.version).where(
                    case_current.c.workspace_id == workspace_id,
                    case_current.c.invoice_id == beta_invoice_id,
                )
            ).one()
            run_count_before_manual = connection.scalar(
                sa.select(sa.func.count())
                .select_from(metadata.tables["recon_run"])
                .where(metadata.tables["recon_run"].c.workspace_id == workspace_id)
            )
        manual_response = client.post(
            "/api/v1/manual-recon",
            json={
                "invoice_id": str(beta_invoice_id),
                "po_ids": [str(beta_po_document_id)],
                "grn_ids": [str(grn_document_id)],
                "run_date": "2026-04-30",
            },
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert manual_response.status_code == 200, manual_response.text
        assert manual_response.json()["invoice_id"] == str(beta_invoice_id)
        assert manual_response.json()["status"] == beta_after_grn.effective_status
        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["recon_run"])
                    .where(metadata.tables["recon_run"].c.workspace_id == workspace_id)
                )
                == run_count_before_manual
            )
            assert (
                connection.scalar(
                    sa.select(case_current.c.version).where(
                        case_current.c.workspace_id == workspace_id,
                        case_current.c.invoice_id == beta_invoice_id,
                    )
                )
                == beta_case_version
            )

        # Case evidence, ETag-protected review and stale-write conflict use the
        # persisted case version. The human decision overlays effective status.
        case_detail = client.get(f"/api/v1/cases/{beta_case_id}")
        assert case_detail.status_code == 200, case_detail.text
        assert case_detail.json()["case"]["invoice_id"] == str(beta_invoice_id)
        assert any(doc["doc_type"] == "invoice" for doc in case_detail.json()["documents"])
        assert case_detail.headers["ETag"] == f'"{beta_case_version}"'
        missing_approval_reason = client.post(
            f"/api/v1/cases/{beta_case_id}/actions",
            json={"action": "approve"},
            headers={
                "If-Match": case_detail.headers["ETag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert missing_approval_reason.status_code == 422
        assert missing_approval_reason.json()["code"] == "approval_reason_required"
        with ingestion_engine.connect() as connection:
            approval_case = connection.execute(
                sa.select(
                    metadata.tables["recon_case"].c.invoice_id,
                    metadata.tables["recon_case"].c.status,
                    metadata.tables["recon_case"].c.reason_codes,
                    metadata.tables["recon_case"].c.input_fingerprint,
                ).where(metadata.tables["recon_case"].c.id == beta_case_id)
            ).one()
        with (
            pytest.raises(IntegrityError) as database_approval_rejection,
            ingestion_engine.begin() as connection,
        ):
            connection.execute(
                sa.insert(metadata.tables["case_action"]).values(
                    id=new_id(),
                    workspace_id=workspace_id,
                    invoice_id=approval_case.invoice_id,
                    case_id=beta_case_id,
                    action="approve",
                    actor_id=user_id,
                    actor_role="manager",
                    computed_status=approval_case.status,
                    reason_codes=approval_case.reason_codes,
                    input_fingerprint=approval_case.input_fingerprint,
                )
            )
        assert (
            database_approval_rejection.value.orig.diag.constraint_name == "case_action_approve_ck"
        )

        def post_review(action_name: str) -> Any:
            return client.post(
                f"/api/v1/cases/{beta_case_id}/actions",
                json={"action": action_name, "reason": f"Synthetic reviewer {action_name}"},
                headers={
                    "If-Match": case_detail.headers["ETag"],
                    "X-CSRF-Token": "synthetic-csrf",
                },
            )

        with ThreadPoolExecutor(max_workers=2) as executor:
            concurrent_reviews = list(executor.map(post_review, ("hold", "reject")))
        assert sorted(response.status_code for response in concurrent_reviews) == [200, 409]
        winning_review = next(
            response for response in concurrent_reviews if response.status_code == 200
        )
        assert winning_review.json()["effective_status"] in {"on_hold", "rejected"}
        with ingestion_engine.connect() as connection:
            trigger_overlay = connection.scalar(sa.select(sa.func.case_overlay(beta_invoice_id)))
            current_overlay = connection.scalar(
                sa.select(case_current.c.overlay).where(
                    case_current.c.workspace_id == workspace_id,
                    case_current.c.invoice_id == beta_invoice_id,
                )
            )
            competing_action_count = connection.scalar(
                sa.select(sa.func.count())
                .select_from(metadata.tables["case_action"])
                .where(
                    metadata.tables["case_action"].c.workspace_id == workspace_id,
                    metadata.tables["case_action"].c.case_id == beta_case_id,
                    metadata.tables["case_action"].c.action.in_(("hold", "reject")),
                )
            )
        assert (
            str(trigger_overlay)
            == str(current_overlay)
            == winning_review.json()["effective_status"]
        )
        assert competing_action_count == 1
        calibration_response = client.get("/api/v1/calibration")
        assert calibration_response.status_code == 200, calibration_response.text
        assert any(
            row["case_id"] == str(beta_case_id) and row["action"] == winning_review.json()["action"]
            for row in calibration_response.json()["overrides"]
        )
        bulk_release_body = {
            "items": [
                {
                    "case_id": str(beta_case_id),
                    "expected_version": winning_review.json()["version"],
                    "action": "release",
                }
            ]
        }
        bulk_release_headers = {
            "Idempotency-Key": "phase4-release-held-case",
            "X-CSRF-Token": "synthetic-csrf",
        }
        released = client.post(
            "/api/v1/cases/bulk-actions",
            json=bulk_release_body,
            headers=bulk_release_headers,
        )
        assert released.status_code == 200, released.text
        assert released.json()["items"][0]["effective_status"] in {"matched", "matched_flags"}
        release_replay = client.post(
            "/api/v1/cases/bulk-actions",
            json=bulk_release_body,
            headers=bulk_release_headers,
        )
        assert release_replay.status_code == 200, release_replay.text
        assert release_replay.json() == released.json()
        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["case_action"])
                    .where(
                        metadata.tables["case_action"].c.workspace_id == workspace_id,
                        metadata.tables["case_action"].c.case_id == beta_case_id,
                        metadata.tables["case_action"].c.action == "release",
                    )
                )
                == 1
            )

        final_kpis = client.get("/api/v1/kpis?from=2026-04-01&to=2026-04-30")
        assert final_kpis.status_code == 200, final_kpis.text
        api_kpis = final_kpis.json()

        # Exercise the export API, worker service, authenticated download, and
        # workbook read-back as one local synthetic end-to-end path.
        app.state.export_service = ExportService(ingestion_engine, storage)
        export_response = client.post(
            "/api/v1/exports",
            json={
                "run_id": str(full_run_id),
                "filter": {"from": "2026-04-01", "to": "2026-04-30"},
            },
            headers={
                "Idempotency-Key": "phase4-beta-export",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert export_response.status_code == 202, export_response.text
        export_id = UUID(export_response.json()["id"])
        export_replay = client.post(
            "/api/v1/exports",
            json={
                "run_id": str(full_run_id),
                "filter": {"from": "2026-04-01", "to": "2026-04-30"},
            },
            headers={
                "Idempotency-Key": "phase4-beta-export",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert export_replay.status_code == 202, export_replay.text
        assert UUID(export_replay.json()["id"]) == export_id
        assert queue.jobs.count(("build_export", export_id)) == 1
        assert queue.jobs[-1] == ("build_export", export_id)
        app.state.export_service.build_export(export_id)
        export_state = client.get(f"/api/v1/exports/{export_id}")
        assert export_state.status_code == 200, export_state.text
        assert export_state.json()["status"] == "ready"
        workflow_response = client.get("/api/v1/workflow-status?from=2026-04-01&to=2026-04-30")
        assert workflow_response.status_code == 200, workflow_response.text
        workflow = workflow_response.json()
        steps = {step["key"]: step for step in workflow["steps"]}
        document_table = metadata.tables["document"]
        with ingestion_engine.connect() as connection:
            unassigned_documents = connection.scalar(
                sa.select(sa.func.count())
                .select_from(document_table)
                .where(
                    document_table.c.workspace_id == workspace_id,
                    document_table.c.withdrawn_at.is_(None),
                    document_table.c.vendor_id.is_(None),
                )
            )
        assert workflow["run_id"] == str(full_run_id)
        assert steps["ingest"]["complete"] is True
        assert steps["map"]["unassigned_documents"] == unassigned_documents
        assert steps["map"]["complete"] is (unassigned_documents == 0)
        assert steps["reconcile"]["complete"] is True
        assert steps["export"]["complete"] is True
        assert workflow["next_step"] == (None if unassigned_documents == 0 else "map")
        assert (
            client.get("/api/v1/workflow-status?from=2026-04-30&to=2026-04-01").status_code == 422
        )
        with ingestion_engine.connect() as connection:
            export_totals = (
                connection.execute(
                    sa.text(
                        """
                    SELECT count(*) AS invoice_count,
                           COALESCE(sum(invoice_total), 0) AS invoiced_total,
                           COALESCE(sum(payable_taxable), 0) AS payable_total,
                           COALESCE(sum(at_risk), 0) AS at_risk_total
                      FROM v_case_current
                     WHERE workspace_id = :workspace_id
                       AND (invoice_date IS NULL
                            OR invoice_date BETWEEN DATE '2026-04-01' AND DATE '2026-04-30')
                    """
                    ),
                    {"workspace_id": workspace_id},
                )
                .mappings()
                .one()
            )
        assert export_state.json()["case_count"] == export_totals["invoice_count"]
        workbook_response = client.get(f"/api/v1/exports/{export_id}/download")
        assert workbook_response.status_code == 200, workbook_response.text
        workbook = load_workbook(
            io.BytesIO(workbook_response.content), read_only=True, data_only=False
        )
        cached_workbook = load_workbook(
            io.BytesIO(workbook_response.content), read_only=True, data_only=True
        )
        try:
            assert workbook.sheetnames == [
                "Summary",
                "Invoice Recon",
                "Line Recon",
                "Exceptions",
                "Vendor Summary",
            ]
            assert workbook["Summary"]["B4"].value == api_kpis["invoice_count"]
            # Home KPI cards equal the Excel Summary (App Flow, Home).
            assert workbook["Summary"]["B6"].value == api_kpis["pending_count"]
            assert workbook["Summary"]["B7"].value == api_kpis["exception_count"]
            assert workbook["Summary"]["B8"].value == api_kpis["unlinked_count"]
            assert workbook["Summary"]["B11"].value.startswith("=SUM(")
            assert cached_workbook["Summary"]["B11"].value == float(export_totals["invoiced_total"])
            assert cached_workbook["Summary"]["B12"].value == float(export_totals["payable_total"])
            assert cached_workbook["Summary"]["B13"].value == float(export_totals["at_risk_total"])
            invoice_row = next(
                row
                for row in workbook["Invoice Recon"].iter_rows(min_row=2, values_only=True)
                if row[2] == "APR-26-INV-901"
            )
            assert invoice_row[16] in {"matched", "matched_flags"}
            assert workbook["Invoice Recon"]["K2"].value.startswith("=")
        finally:
            workbook.close()
            cached_workbook.close()

        failed_export_response = client.post(
            "/api/v1/exports",
            json={"run_id": str(full_run_id), "scope": "all"},
            headers={
                "Idempotency-Key": "phase4-export-retry",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert failed_export_response.status_code == 202, failed_export_response.text
        failed_export_id = UUID(failed_export_response.json()["id"])
        original_put_fileobj = storage.put_fileobj

        def fail_object_write(*args: Any, **kwargs: Any) -> None:
            del args, kwargs
            raise OSError("synthetic object-store interruption")

        monkeypatch.setattr(storage, "put_fileobj", fail_object_write)
        with pytest.raises(OSError, match="synthetic object-store interruption"):
            app.state.export_service.build_export(failed_export_id)
        monkeypatch.setattr(storage, "put_fileobj", original_put_fileobj)
        failed_export_state = client.get(f"/api/v1/exports/{failed_export_id}")
        assert failed_export_state.status_code == 200
        assert failed_export_state.json()["status"] == "failed"
        assert failed_export_state.json()["error_code"] == "EXPORT_BUILD_FAILED"
        retried_export = client.post(
            f"/api/v1/exports/{failed_export_id}/retry",
            headers={
                "If-Match": failed_export_state.headers["ETag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert retried_export.status_code == 202, retried_export.text
        assert queue.jobs.count(("build_export", failed_export_id)) == 2
        app.state.export_service.build_export(failed_export_id)
        assert client.get(f"/api/v1/exports/{failed_export_id}").json()["status"] == "ready"

        vendor_documents_response = client.get(
            f"/api/v1/vendors/{resolved_vendor_id}/documents?doc_type=invoice"
        )
        assert vendor_documents_response.status_code == 200, vendor_documents_response.text
        listed_linked_invoice = next(
            row for row in vendor_documents_response.json() if row["id"] == str(linked_invoice_id)
        )
        assert listed_linked_invoice["links"][0]["po_no"] == "PO-2026-001"
        assert listed_linked_invoice["links"][0]["version_no"] == 1

        bulk_body = {"vendor_id": str(resolved_vendor_id), "threshold": "0.95"}
        bulk_headers = {
            "Idempotency-Key": "phase4-bulk-confirm-once",
            "X-CSRF-Token": "synthetic-csrf",
        }
        bulk_confirm = client.post(
            "/api/v1/links/bulk-confirm", json=bulk_body, headers=bulk_headers
        )
        assert bulk_confirm.status_code == 200, bulk_confirm.text
        assert bulk_confirm.json()["confirmed"] >= 1
        bulk_replay = client.post(
            "/api/v1/links/bulk-confirm", json=bulk_body, headers=bulk_headers
        )
        assert bulk_replay.status_code == 200, bulk_replay.text
        assert bulk_replay.headers["Idempotency-Replayed"] == "true"
        assert bulk_replay.json() == bulk_confirm.json()

        current_invoice_links = client.get(f"/api/v1/links?source_document_id={linked_invoice_id}")
        assert current_invoice_links.status_code == 200, current_invoice_links.text
        current_link = current_invoice_links.json()[0]
        assert current_link["state"] == "confirmed"
        stale_edit = client.patch(
            f"/api/v1/links/{current_link['id']}",
            json={"action": "unlink"},
            headers={"If-Match": '"1"', "X-CSRF-Token": "synthetic-csrf"},
        )
        assert stale_edit.status_code == 409, stale_edit.text
        unlinked = client.patch(
            f"/api/v1/links/{current_link['id']}",
            json={"action": "unlink"},
            headers={
                "If-Match": f'"{current_link["version_no"]}"',
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert unlinked.status_code == 200, unlinked.text
        assert unlinked.json()["state"] == "rejected"
        mapping_service.resolve_workspace(workspace_id)
        persisted_unlink = client.get(f"/api/v1/links?source_document_id={linked_invoice_id}")
        assert persisted_unlink.json()[0]["state"] == "rejected"

        # Model an unresolved document and exercise the user-controlled assignment path.
        with ingestion_engine.begin() as connection:
            connection.execute(
                sa.update(metadata.tables["document"])
                .where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.id == document_id,
                )
                .values(vendor_id=None, vendor_locked=False, vendor_rule="none")
            )
        unassigned_rows = client.get("/api/v1/vendors/unassigned/documents?doc_type=invoice")
        assert unassigned_rows.status_code == 200, unassigned_rows.text
        assert any(row["id"] == str(document_id) for row in unassigned_rows.json())
        assignment = client.post(
            f"/api/v1/documents/{document_id}/assign-vendor",
            json={"vendor_id": str(resolved_vendor_id)},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert assignment.status_code == 200, assignment.text
        assert assignment.json()["vendor_locked"] is True

        # A same-PAN human merge is soft: source rows remain attached to the old ID.
        same_pan_vendor_id = new_id()
        suggestion_id = new_id()
        vendor_table = metadata.tables["vendor"]
        with ingestion_engine.begin() as connection:
            connection.execute(
                sa.insert(vendor_table).values(
                    id=same_pan_vendor_id,
                    workspace_id=workspace_id,
                    canonical_name="Acme Components Legacy",
                    normalized_name="acme components legacy",
                    pan="AAPFU0939F",
                    anchor="pan",
                    created_by_rule="manual",
                )
            )
            vendor_a_id, vendor_b_id = sorted((same_pan_vendor_id, resolved_vendor_id), key=str)
            connection.execute(
                sa.insert(metadata.tables["vendor_merge_suggestion"]).values(
                    id=suggestion_id,
                    workspace_id=workspace_id,
                    vendor_a_id=vendor_a_id,
                    vendor_b_id=vendor_b_id,
                    score=Decimal("0.87"),
                    state="pending",
                )
            )
            connection.execute(
                sa.update(metadata.tables["document"])
                .where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.id == document_id,
                )
                .values(vendor_id=same_pan_vendor_id, vendor_locked=True, vendor_rule="manual")
            )
        suggestions = client.get("/api/v1/vendor-merge-suggestions")
        assert suggestions.status_code == 200, suggestions.text
        suggestion = next(row for row in suggestions.json() if row["id"] == str(suggestion_id))
        assert suggestion["pan_conflict"] is False
        accepted_merge = client.post(
            f"/api/v1/vendor-merge-suggestions/{suggestion_id}/decision",
            json={"decision": "accept", "survivor_vendor_id": str(resolved_vendor_id)},
            headers={"If-Match": '"1"', "X-CSRF-Token": "synthetic-csrf"},
        )
        assert accepted_merge.status_code == 200, accepted_merge.text
        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(vendor_table.c.merged_into).where(
                        vendor_table.c.workspace_id == workspace_id,
                        vendor_table.c.id == same_pan_vendor_id,
                    )
                )
                == resolved_vendor_id
            )
            assert (
                connection.scalar(
                    sa.select(metadata.tables["document"].c.vendor_id).where(
                        metadata.tables["document"].c.workspace_id == workspace_id,
                        metadata.tables["document"].c.id == document_id,
                    )
                )
                == same_pan_vendor_id
            )

        alias_added = client.post(
            "/api/v1/vendor-aliases",
            json={"vendor_id": str(resolved_vendor_id), "alias": "Acme Components West"},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert alias_added.status_code == 200, alias_added.text
        aliases = client.get(f"/api/v1/vendor-aliases?vendor_id={resolved_vendor_id}")
        assert any(row["id"] == alias_added.json()["id"] for row in aliases.json())
        alias_deleted = client.delete(
            f"/api/v1/vendor-aliases/{alias_added.json()['id']}",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert alias_deleted.status_code == 200, alias_deleted.text

        different_pan_vendor_id = new_id()
        with ingestion_engine.begin() as connection:
            connection.execute(
                sa.insert(vendor_table).values(
                    id=different_pan_vendor_id,
                    workspace_id=workspace_id,
                    canonical_name="Different PAN Vendor",
                    normalized_name="different pan vendor",
                    pan="BBBBB1234B",
                    anchor="pan",
                    created_by_rule="manual",
                )
            )
        cross_pan_merge = client.post(
            f"/api/v1/vendors/{different_pan_vendor_id}/merge",
            json={"survivor_vendor_id": str(resolved_vendor_id)},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert cross_pan_merge.status_code == 422, cross_pan_merge.text
        assert cross_pan_merge.json()["code"] == "invalid_vendor_merge"

        # Upload the no-reference invoice first, then its PO and GRN, to verify
        # that later evidence resolves an earlier Unlinked record safely.
        upload_and_parse(
            "invoice",
            "beta-invoice-no-po-reference.csv",
            _CONTENT_INVOICE_NO_PO_REF_CSV,
            124,
            mapping={
                "PO Ref": "po_refs",
                "Invoice No": "doc_no",
                "Invoice Date": "doc_date",
            },
        )
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service.resolve_workspace(workspace_id)
        with ingestion_engine.connect() as connection:
            no_ref_invoice_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_no == "APR-26-INV-903",
                )
            )
        assert no_ref_invoice_id is not None
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service.resolve_workspace(workspace_id)
        no_ref_links = client.get(f"/api/v1/links?source_document_id={no_ref_invoice_id}")
        assert no_ref_links.status_code == 200, no_ref_links.text
        assert no_ref_links.json() == []

        upload_and_parse("po", "content-purchase-order.csv", _CONTENT_PO_CSV, 126)
        vendor_resolution.resolve_workspace(workspace_id)
        content_po_mapping = mapping_service.resolve_workspace(workspace_id)
        assert content_po_mapping.linked >= 1
        po_only_links = client.get(f"/api/v1/links?source_document_id={no_ref_invoice_id}")
        assert po_only_links.status_code == 200, po_only_links.text
        assert any(
            row["state"] == "suggested" and row["method"] == "content"
            for row in po_only_links.json()
        )

        upload_and_parse(
            "grn",
            "content-receipt.csv",
            _CONTENT_GRN_CSV,
            128,
            mapping={"PO Ref": "po_refs", "GRN No": "doc_no", "GRN Date": "doc_date"},
        )
        vendor_resolution.resolve_workspace(workspace_id)
        content_grn_mapping = mapping_service.resolve_workspace(workspace_id)
        assert content_grn_mapping.linked >= 1
        no_ref_links = client.get(f"/api/v1/links?source_document_id={no_ref_invoice_id}")
        assert no_ref_links.status_code == 200, no_ref_links.text
        content_suggestion = next(
            row
            for row in no_ref_links.json()
            if row["state"] == "suggested" and row["method"] == "content"
        )
        assert Decimal(str(content_suggestion["confidence"])) <= Decimal("0.9000")
        human_confirmation = client.patch(
            f"/api/v1/links/{content_suggestion['id']}",
            json={"action": "confirm"},
            headers={
                "If-Match": f'"{content_suggestion["version_no"]}"',
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert human_confirmation.status_code == 200, human_confirmation.text
        content_calibration = client.get("/api/v1/calibration")
        assert content_calibration.status_code == 200, content_calibration.text
        calibrated_content_decision = next(
            row
            for row in content_calibration.json()["link_decisions"]
            if row["decision_link_id"] == human_confirmation.json()["id"]
        )
        assert calibrated_content_decision["outcome"] == "confirmed"
        assert calibrated_content_decision["method"] == "content"
        assert calibrated_content_decision["suggested_method"] == "content"

        # Compare the same semantic invoice/PO/GRN shape with the reference
        # documents uploaded first. Final business results should agree across
        # these two ingestion orders after an operator confirms the suggestion.
        second_po_csv = _CONTENT_PO_CSV.replace(b"PO-2026-778", b"PO-2026-779").replace(
            b"Widget C Compact", b"Hydraulic Valve Assembly X9"
        )
        second_grn_csv = (
            _CONTENT_GRN_CSV.replace(b"PO-2026-778", b"PO-2026-779")
            .replace(b"GRN-2026-778", b"GRN-2026-779")
            .replace(b"Widget C Compact", b"Hydraulic Valve Assembly X9")
        )
        second_invoice_csv = _CONTENT_INVOICE_NO_PO_REF_CSV.replace(
            b"APR-26-INV-903", b"BETA-26-INV-904"
        ).replace(b"Widget C Compact", b"Hydraulic Valve Assembly X9")
        upload_and_parse("po", "ordered-first-po.csv", second_po_csv, 202)
        upload_and_parse(
            "grn",
            "ordered-first-grn.csv",
            second_grn_csv,
            204,
            mapping={"PO Ref": "po_refs", "GRN No": "doc_no", "GRN Date": "doc_date"},
        )
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service.resolve_workspace(workspace_id)
        upload_and_parse(
            "invoice",
            "ordered-first-invoice.csv",
            second_invoice_csv,
            206,
            mapping={
                "PO Ref": "po_refs",
                "Invoice No": "doc_no",
                "Invoice Date": "doc_date",
            },
        )
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service.resolve_workspace(workspace_id)
        with ingestion_engine.connect() as connection:
            ordered_first_invoice_id = connection.scalar(
                sa.select(metadata.tables["document"].c.id).where(
                    metadata.tables["document"].c.workspace_id == workspace_id,
                    metadata.tables["document"].c.doc_no == "BETA-26-INV-904",
                )
            )
        assert ordered_first_invoice_id is not None
        ordered_first_links = client.get(
            f"/api/v1/links?source_document_id={ordered_first_invoice_id}"
        )
        assert ordered_first_links.status_code == 200, ordered_first_links.text
        ordered_first_suggestion = next(
            row
            for row in ordered_first_links.json()
            if row["state"] == "suggested" and row["method"] == "content"
        )
        ordered_first_confirmed = client.patch(
            f"/api/v1/links/{ordered_first_suggestion['id']}",
            json={"action": "confirm"},
            headers={
                "If-Match": f'"{ordered_first_suggestion["version_no"]}"',
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert ordered_first_confirmed.status_code == 200, ordered_first_confirmed.text
        order_check_run_response = client.post(
            "/api/v1/runs",
            json={"run_date": "2026-04-30"},
            headers={
                "Idempotency-Key": "replay-test-run-order-check",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert order_check_run_response.status_code == 202, order_check_run_response.text
        order_check_run_id = UUID(order_check_run_response.json()["id"])
        app.state.reconciliation_service.evaluate_run(order_check_run_id)
        with ingestion_engine.connect() as connection:
            order_results = (
                connection.execute(
                    sa.select(
                        metadata.tables["recon_case"].c.invoice_id,
                        metadata.tables["recon_case"].c.status,
                        metadata.tables["recon_case"].c.reason_codes,
                        metadata.tables["recon_case"].c.invoice_taxable,
                        metadata.tables["recon_case"].c.invoice_total,
                        metadata.tables["recon_case"].c.expected_taxable,
                        metadata.tables["recon_case"].c.variance,
                        metadata.tables["recon_case"].c.payable_taxable,
                        metadata.tables["recon_case"].c.at_risk,
                        metadata.tables["recon_case"].c.at_risk_gst,
                        metadata.tables["recon_case"].c.blocked_value,
                    )
                    .where(
                        metadata.tables["recon_case"].c.workspace_id == workspace_id,
                        metadata.tables["recon_case"].c.run_id == order_check_run_id,
                        metadata.tables["recon_case"].c.invoice_id.in_(
                            (no_ref_invoice_id, ordered_first_invoice_id)
                        ),
                    )
                    .order_by(metadata.tables["recon_case"].c.invoice_id)
                )
                .mappings()
                .all()
            )
        assert len(order_results) == 2
        business_fields = (
            "invoice_taxable",
            "invoice_total",
            "expected_taxable",
            "variance",
            "payable_taxable",
            "at_risk",
            "at_risk_gst",
            "blocked_value",
        )
        assert tuple(order_results[0][key] for key in business_fields) == tuple(
            order_results[1][key] for key in business_fields
        )
        first_reasons = set(order_results[0]["reason_codes"])
        second_reasons = set(order_results[1]["reason_codes"])
        assert (ReasonCode.DUP_SUSPECTED.value in first_reasons) != (
            ReasonCode.DUP_SUSPECTED.value in second_reasons
        )

        repeated = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=_CSV,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "another-name.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert repeated.status_code == 200, repeated.text
        assert repeated.json()["duplicate"] is True
        assert repeated.json()["file_id"] == str(file_id)

        with ingestion_engine.connect() as connection:
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["document"])
                    .where(metadata.tables["document"].c.workspace_id == workspace_id)
                )
                == 13
            )
            assert (
                connection.scalar(
                    sa.select(sa.func.count())
                    .select_from(metadata.tables["document_version"])
                    .where(metadata.tables["document_version"].c.workspace_id == workspace_id)
                )
                == 13
            )
            assert (
                connection.scalar(
                    sa.select(metadata.tables["source_file"].c.reupload_skips).where(
                        metadata.tables["source_file"].c.id == file_id
                    )
                )
                == 1
            )

        deleted = client.delete(
            f"/api/v1/files/{file_id}", headers={"X-CSRF-Token": "synthetic-csrf"}
        )
        assert deleted.status_code == 204
        assert client.get(f"/api/v1/documents/{document_id}").status_code == 404

        reactivated = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=_CSV,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "reactivated-invoice.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert reactivated.status_code == 201, reactivated.text
        assert reactivated.json()["reactivated"] is True
        assert reactivated.json()["file_id"] == str(file_id)
        process = client.post(
            f"/api/v1/batches/{batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(file_id, job_id=3)
        service.parse_file(file_id, job_id=4)
        assert client.get(f"/api/v1/documents/{document_id}").status_code == 200

        suggested_batch_response = client.post(
            "/api/v1/batches",
            json={},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert suggested_batch_response.status_code == 201, suggested_batch_response.text
        suggested_batch_id = UUID(suggested_batch_response.json()["id"])
        suggested_csv = _CSV.replace(b"INV-2026-001", b"INV-2026-003").replace(
            b"01/04/2026", b"05/04/2026"
        )
        suggested_upload = client.post(
            f"/api/v1/files?batch_id={suggested_batch_id}&zone=invoice",
            content=suggested_csv,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "invoice-003.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert suggested_upload.status_code == 201, suggested_upload.text
        suggested_file_id = UUID(suggested_upload.json()["file_id"])
        process = client.post(
            f"/api/v1/batches/{suggested_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(suggested_file_id, job_id=8)
        service.parse_file(suggested_file_id, job_id=9)
        suggested_summary = client.get(f"/api/v1/batches/{suggested_batch_id}/summary")
        assert suggested_summary.status_code == 200
        assert suggested_summary.json()["label"] == "Apr 2026"
        assert suggested_summary.json()["label_is_suggested"] is True

        mismatch_csv = _CSV.replace(b"INV-2026-001", b"INV-2026-004")
        mismatch_upload = client.post(
            f"/api/v1/files?batch_id={suggested_batch_id}&zone=po",
            content=mismatch_csv,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "mis-zoned.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert mismatch_upload.status_code == 201, mismatch_upload.text
        mismatch_file_id = UUID(mismatch_upload.json()["file_id"])
        process = client.post(
            f"/api/v1/batches/{suggested_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(mismatch_file_id, job_id=10)
        mismatch_state = client.get(f"/api/v1/files/{mismatch_file_id}")
        assert mismatch_state.status_code == 200
        assert mismatch_state.json()["mismatch_state"] == "suggested"
        resolved = client.patch(
            f"/api/v1/files/{mismatch_file_id}",
            json={"decision": "accept_suggested"},
            headers={
                "If-Match": mismatch_state.headers["ETag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert resolved.status_code == 200, resolved.text
        assert resolved.json()["zone"] == "invoice"
        assert resolved.json()["mismatch_state"] == "accepted"
        process = client.post(
            f"/api/v1/batches/{suggested_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(mismatch_file_id, job_id=11)
        service.parse_file(mismatch_file_id, job_id=12)
        assert client.get(f"/api/v1/files/{mismatch_file_id}").json()["status"] == "parsed"
        errors_response = client.get(f"/api/v1/files/{mismatch_file_id}/errors")
        assert errors_response.status_code == 200
        error_rows = [json.loads(line) for line in errors_response.text.splitlines()]
        mismatch_warnings = [row for row in error_rows if row.get("error_code") == "TYPE_MISMATCH"]
        assert mismatch_warnings
        assert all(row.get("resolved_at") is not None for row in mismatch_warnings)

        archive_body = _zip_bytes(
            {
                "invoices/invoice-002.csv": _CSV.replace(b"INV-2026-001", b"INV-2026-002"),
                "nested.zip": _zip_bytes({"inside.csv": _CSV}),
                "__MACOSX/._metadata": b"ignored",
            }
        )
        archive_upload = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=archive_body,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "month.zip",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert archive_upload.status_code == 201, archive_upload.text
        archive_id = UUID(archive_upload.json()["file_id"])
        process = client.post(
            f"/api/v1/batches/{batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(archive_id, job_id=5)
        with ingestion_engine.connect() as connection:
            archive_children = (
                connection.execute(
                    sa.select(metadata.tables["source_file"])
                    .where(metadata.tables["source_file"].c.parent_file_id == archive_id)
                    .order_by(metadata.tables["source_file"].c.archive_path)
                )
                .mappings()
                .all()
            )
        assert len(archive_children) == 2
        csv_child = next(row for row in archive_children if row["format"] == "csv")
        nested_child = next(row for row in archive_children if row["format"] == "zip")
        assert nested_child["status"] == "failed"
        assert nested_child["error_code"] == "NESTED_ARCHIVE_REJECTED"
        service.parse_file(UUID(str(csv_child["id"])), job_id=6)
        archive_status = client.get(f"/api/v1/files/{archive_id}").json()
        assert archive_status["status"] == "partial"
        assert archive_status["docs_found"] == 1

        traversal = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone=invoice",
            content=_zip_bytes({"../escape.csv": _CSV}),
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "hostile.zip",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert traversal.status_code == 201, traversal.text
        traversal_id = UUID(traversal.json()["file_id"])
        process = client.post(
            f"/api/v1/batches/{batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert process.status_code == 202, process.text
        service.ingest_file(traversal_id, job_id=7)
        failed = client.get(f"/api/v1/files/{traversal_id}").json()
        assert failed["status"] == "failed"
        assert failed["error_code"] == "ZIP_PATH_TRAVERSAL"
        error_dashboard = client.get("/api/v1/admin/errors")
        assert error_dashboard.status_code == 200, error_dashboard.text
        assert any(
            row["id"] == str(traversal_id) and row["error_code"] == "ZIP_PATH_TRAVERSAL"
            for row in error_dashboard.json()["retry_actions"]
        )
        assert any(
            row["error_code"] == "ZIP_PATH_TRAVERSAL"
            for row in error_dashboard.json()["failed_by_reason_and_day"]
        )
        retried = client.post(
            f"/api/v1/files/{traversal_id}/retry",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert retried.status_code == 202, retried.text
        assert retried.json()["status"] == "queued"
        assert queue.jobs[-1] == ("ingest_file", traversal_id)
        service.ingest_file(traversal_id, job_id=8)
        failed_again = client.get(f"/api/v1/files/{traversal_id}").json()
        assert failed_again["status"] == "failed"
        assert failed_again["error_code"] == "ZIP_PATH_TRAVERSAL"

        pdf_batch_response = client.post(
            "/api/v1/batches",
            json={"label": "Synthetic PDF batch"},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert pdf_batch_response.status_code == 201, pdf_batch_response.text
        pdf_batch_id = UUID(pdf_batch_response.json()["id"])
        pdf_content = _pdf_bytes(
            "Invoice No: SYN-PDF-001",
            "Invoice Date: 2026-04-02",
            "Supplier: Synthetic PDF Vendor",
            "Vendor GSTIN: 27AAPFU0939F1ZV",
            "PO Reference: SYN-PO-0142",
            "Quantity: 1",
            "Rate: 100",
            "Taxable Value: 100",
            "Grand Total: 100",
        )
        pdf_upload = client.post(
            f"/api/v1/files?batch_id={pdf_batch_id}&zone=invoice",
            content=pdf_content,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "synthetic-invoice.pdf",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert pdf_upload.status_code == 201, pdf_upload.text
        pdf_file_id = UUID(pdf_upload.json()["file_id"])
        queued_pdf = client.post(
            f"/api/v1/batches/{pdf_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert queued_pdf.status_code == 202, queued_pdf.text
        service.ingest_file(pdf_file_id, job_id=15)
        service.parse_file(pdf_file_id, job_id=16)
        pdf_file = client.get(f"/api/v1/files/{pdf_file_id}")
        assert pdf_file.status_code == 200
        assert pdf_file.json()["status"] == "parsed", pdf_file.json()

        pdf_document = client.get(f"/api/v1/documents?batch_id={pdf_batch_id}")
        assert pdf_document.status_code == 200
        assert len(pdf_document.json()) == 1
        pdf_document_id = UUID(pdf_document.json()[0]["id"])
        pdf_details = client.get(f"/api/v1/documents/{pdf_document_id}")
        assert pdf_details.status_code == 200
        doc_number = next(
            item for item in pdf_details.json()["field_evidence"] if item["field_name"] == "doc_no"
        )
        assert doc_number["page"] == 1
        assert doc_number["bbox"] is not None
        with ingestion_engine.connect() as connection:
            version_id = connection.scalar(
                sa.select(metadata.tables["document"].c.current_version_id).where(
                    metadata.tables["document"].c.id == pdf_document_id
                )
            )
            text_row = (
                connection.execute(
                    sa.select(metadata.tables["document_text"]).where(
                        metadata.tables["document_text"].c.document_version_id == version_id,
                        metadata.tables["document_text"].c.page == 1,
                    )
                )
                .mappings()
                .one()
            )
        assert "SYN-PDF-001" in text_row["content"]
        assert text_row["char_count"] == len(text_row["content"])

        source_pdf_details = client.get(f"/api/v1/documents/{pdf_document_id}")
        original_pdf_version = UUID(source_pdf_details.json()["version"]["id"])
        resolve_jobs_before_correction = queue.jobs.count(("resolve_vendors", pdf_file_id))
        corrected_pdf = client.put(
            f"/api/v1/documents/{pdf_document_id}/correction",
            json={
                "doc_no": "SYN-PDF-001-REVIEWED",
                "lines": [
                    {
                        "line_no": 1,
                        "description": "Reviewer-entered synthetic line",
                        "quantity": "1",
                        "rate": "100",
                        "amount": "100",
                    }
                ],
                "verified": True,
                "reason": "Synthetic integration check of human review.",
            },
            headers={
                "If-Match": source_pdf_details.headers["etag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert corrected_pdf.status_code == 200, corrected_pdf.text
        assert queue.jobs.count(("resolve_vendors", pdf_file_id)) == (
            resolve_jobs_before_correction + 1
        )
        corrected_details = corrected_pdf.json()
        corrected_version_id = UUID(corrected_details["version"]["id"])
        assert corrected_version_id != original_pdf_version
        assert corrected_details["version"]["extraction_method"] == "manual"
        assert corrected_details["version"]["needs_verification"] is False
        assert corrected_details["document"]["doc_no"] == "SYN-PDF-001-REVIEWED"
        assert corrected_details["lines"][0]["description"] == "Reviewer-entered synthetic line"
        assert any(
            item["field_name"] == "description" and item["source"] == "manual"
            for item in corrected_details["field_evidence"]
        )
        with ingestion_engine.connect() as connection:
            copied_page = connection.execute(
                sa.select(metadata.tables["document_text"].c.content).where(
                    metadata.tables["document_text"].c.workspace_id == workspace_id,
                    metadata.tables["document_text"].c.document_version_id == corrected_version_id,
                    metadata.tables["document_text"].c.page == 1,
                )
            ).scalar_one()
        assert "SYN-PDF-001" in copied_page

        stale_review = client.put(
            f"/api/v1/documents/{pdf_document_id}/correction",
            json={"verified": True},
            headers={
                "If-Match": source_pdf_details.headers["etag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert stale_review.status_code == 409

        shell_content = _scanned_pdf_bytes()
        shell_upload = client.post(
            f"/api/v1/files?batch_id={pdf_batch_id}&zone=invoice",
            content=shell_content,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "synthetic-scan.pdf",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert shell_upload.status_code == 201, shell_upload.text
        shell_file_id = UUID(shell_upload.json()["file_id"])
        queued_shell = client.post(
            f"/api/v1/batches/{pdf_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert queued_shell.status_code == 202, queued_shell.text
        service.ingest_file(shell_file_id, job_id=17)
        service.parse_file(shell_file_id, job_id=18)
        shell_file = client.get(f"/api/v1/files/{shell_file_id}")
        assert shell_file.json()["status"] == "partial"
        assert shell_file.json()["retry_available"] is True
        shell_errors = client.get(f"/api/v1/files/{shell_file_id}/errors")
        assert shell_errors.status_code == 200
        shell_warning_rows = [json.loads(line) for line in shell_errors.text.splitlines()]
        assert any(row["error_code"] == "NO_TEXT_LAYER" for row in shell_warning_rows)
        retry_shell = client.post(
            f"/api/v1/files/{shell_file_id}/retry",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert retry_shell.status_code == 202, retry_shell.text
        assert retry_shell.json()["status"] == "queued"
        with ingestion_engine.connect() as connection:
            shell_document = (
                connection.execute(
                    sa.select(metadata.tables["document"])
                    .where(
                        metadata.tables["document"].c.workspace_id == workspace_id,
                    )
                    .join(
                        metadata.tables["document_version"],
                        sa.and_(
                            metadata.tables["document_version"].c.workspace_id
                            == metadata.tables["document"].c.workspace_id,
                            metadata.tables["document_version"].c.document_id
                            == metadata.tables["document"].c.id,
                            metadata.tables["document_version"].c.id
                            == metadata.tables["document"].c.current_version_id,
                        ),
                    )
                    .where(metadata.tables["document_version"].c.source_file_id == shell_file_id)
                )
                .mappings()
                .one()
            )
            shell_version = (
                connection.execute(
                    sa.select(metadata.tables["document_version"]).where(
                        metadata.tables["document_version"].c.id
                        == shell_document["current_version_id"]
                    )
                )
                .mappings()
                .one()
            )
            shell_page_hashes = connection.scalar(
                sa.select(metadata.tables["document_identity"].c.page_hashes).where(
                    metadata.tables["document_identity"].c.workspace_id == workspace_id,
                    metadata.tables["document_identity"].c.document_version_id
                    == shell_version["id"],
                )
            )
        expected_shell_hashes = page_image_hashes_isolated(shell_content, (1,))
        assert shell_version["needs_verification"] is True
        assert "NO_TEXT_LAYER" in shell_version["verification_reasons"]
        assert expected_shell_hashes is not None
        assert shell_page_hashes == [item.value for item in expected_shell_hashes]

        unreadable_pdf = b"%PDF-1.4\nThis is not a readable PDF body."
        unreadable_upload = client.post(
            f"/api/v1/files?batch_id={pdf_batch_id}&zone=invoice",
            content=unreadable_pdf,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "unreadable-invoice.pdf",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert unreadable_upload.status_code == 201, unreadable_upload.text
        unreadable_file_id = UUID(unreadable_upload.json()["file_id"])
        queued_unreadable = client.post(
            f"/api/v1/batches/{pdf_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert queued_unreadable.status_code == 202, queued_unreadable.text
        service.ingest_file(unreadable_file_id, job_id=31)
        service.parse_file(unreadable_file_id, job_id=32)
        source_documents = client.get(
            f"/api/v1/documents?batch_id={pdf_batch_id}&source_file_id={unreadable_file_id}"
        )
        assert source_documents.json() == []
        manual_entry = client.post(
            f"/api/v1/files/{unreadable_file_id}/manual-invoice",
            json={
                "doc_no": "MANUAL-SYN-001",
                "vendor_name": "Synthetic Manual Vendor",
                "po_refs": ["PO/MANUAL/SYN/001"],
                "grand_total": "118.00",
                "lines": [
                    {
                        "line_no": 1,
                        "description": "Manual entry test item",
                        "quantity": "1",
                        "rate": "100.00",
                        "amount": "100.00",
                    }
                ],
                "verified": True,
                "reason": "Synthetic manual-entry integration test.",
            },
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert manual_entry.status_code == 201, manual_entry.text
        manual_details = manual_entry.json()
        assert manual_details["version"]["extraction_method"] == "manual"
        assert manual_details["version"]["needs_verification"] is False
        assert manual_details["version"]["doc_no"] == "MANUAL-SYN-001"
        assert manual_details["lines"][0]["description"] == "Manual entry test item"
        assert all(item["source"] == "manual" for item in manual_details["field_evidence"])
        manual_doc_no_evidence = next(
            item for item in manual_details["field_evidence"] if item["field_name"] == "doc_no"
        )
        assert manual_doc_no_evidence["origin"] == FieldOrigin.MANUAL.value
        assert manual_doc_no_evidence["confidence"] is None
        assert manual_doc_no_evidence["verification"] == FieldVerification.USER.value
        assert manual_doc_no_evidence["verified_by"] is not None
        assert manual_doc_no_evidence["verified_at"] is not None
        manual_po_evidence = next(
            item for item in manual_details["field_evidence"] if item["field_name"] == "po_refs"
        )
        assert manual_po_evidence["origin"] == FieldOrigin.MANUAL.value
        assert manual_po_evidence["confidence"] is None
        assert manual_po_evidence["verification"] == FieldVerification.USER.value

        from reconbench_app.ingestion import _preserve_user_verified_references

        extracted_reference = FieldValue(
            value=("PO-MANUAL-SYN-001",),
            confidence=Decimal("0.99"),
            source=FieldSource.RULE,
            evidence=FieldEvidence(snippet="Synthetic OCR text"),
            origin=FieldOrigin.OCR,
        )
        extracted_document = Document(
            header=InvoiceHeader(doc_type=DocType.INVOICE, po_refs=extracted_reference)
        )
        with ingestion_engine.begin() as connection:
            preserved_document, conflict = _preserve_user_verified_references(
                connection,
                workspace_id=workspace_id,
                version_id=UUID(manual_details["version"]["id"]),
                document=extracted_document,
            )
        assert not conflict
        preserved_reference = preserved_document.header.po_refs
        assert preserved_reference is not None
        assert preserved_reference.value == ("PO/MANUAL/SYN/001",)
        assert preserved_reference.origin is FieldOrigin.MANUAL
        assert preserved_reference.confidence is None
        assert preserved_reference.verification is FieldVerification.USER
        updated_unreadable = client.get(f"/api/v1/files/{unreadable_file_id}").json()
        assert updated_unreadable["status"] == "parsed"
        assert updated_unreadable["docs_found"] == 1

        mapping_batch_response = client.post(
            "/api/v1/batches",
            json={"label": "Synthetic CSV mapping batch"},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert mapping_batch_response.status_code == 201, mapping_batch_response.text
        mapping_batch_id = UUID(mapping_batch_response.json()["id"])
        unmapped_csv = (
            b"Field Alpha,Field Beta,Field Gamma,Field Delta,Field Epsilon,Field Zeta,"
            b"Field Eta,Field Theta,Field Iota\n"
            b"SYN-MAPPED-001,2026-04-03,Mapping Vendor,Widget,1,100,100,100,100\n"
        )
        unmapped_upload = client.post(
            f"/api/v1/files?batch_id={mapping_batch_id}&zone=invoice",
            content=unmapped_csv,
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": "custom-layout.csv",
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert unmapped_upload.status_code == 201, unmapped_upload.text
        unmapped_file_id = UUID(unmapped_upload.json()["file_id"])
        queue_unmapped = client.post(
            f"/api/v1/batches/{mapping_batch_id}/process",
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert queue_unmapped.status_code == 202, queue_unmapped.text
        service.ingest_file(unmapped_file_id, job_id=19)
        service.parse_file(unmapped_file_id, job_id=20)
        unmapped_state = client.get(f"/api/v1/files/{unmapped_file_id}")
        assert unmapped_state.status_code == 200
        assert unmapped_state.json()["error_code"] == "CSV_MAPPING_REQUIRED"
        saved_mapping = client.post(
            f"/api/v1/files/{unmapped_file_id}/column-mapping",
            json={
                "column_mapping": {
                    "Field Alpha": "doc_no",
                    "Field Beta": "doc_date",
                    "Field Gamma": "vendor_name",
                    "Field Delta": "description",
                    "Field Epsilon": "quantity",
                    "Field Zeta": "rate",
                    "Field Eta": "amount",
                    "Field Theta": "taxable_total",
                    "Field Iota": "grand_total",
                }
            },
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert saved_mapping.status_code == 202, saved_mapping.text
        service.ingest_file(unmapped_file_id, job_id=21)
        service.parse_file(unmapped_file_id, job_id=22)
        mapped_state = client.get(f"/api/v1/files/{unmapped_file_id}")
        assert mapped_state.status_code == 200
        assert mapped_state.json()["status"] == "parsed", mapped_state.json()
        mapped_documents = client.get(f"/api/v1/documents?batch_id={mapping_batch_id}")
        assert mapped_documents.status_code == 200
        assert len(mapped_documents.json()) == 1
        assert mapped_documents.json()[0]["doc_no"] == "SYN-MAPPED-001"


def test_po_and_grn_corrections_save_versions_and_refresh_reconciliation(
    ingestion_engine: Engine,
    tmp_path: Any,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("RB_AI_EXTERNAL_CALLS", "disabled")
    workspace_id, user_id = _seed_workspace(ingestion_engine)
    fake_s3 = _FakeObjectStoreClient()
    monkeypatch.setattr(boto3, "client", lambda *args, **kwargs: fake_s3)
    storage = LocalObjectStore(tmp_path / "objects")
    settings = Settings(
        _env_file=None,
        environment="ci",
        database_url=ingestion_engine.url.render_as_string(hide_password=False),
        s3_endpoint_url="http://127.0.0.1:1",
        s3_region="us-east-1",
        s3_bucket="synthetic-corrections",
        s3_api_access_key="synthetic-access",
        s3_api_secret_key="synthetic-secret",
        public_base_url="https://example.invalid",
        argon2_time_cost=1,
        argon2_memory_cost_kib=8192,
        argon2_parallelism=1,
    )
    app = create_app(settings)
    queue = _FakeQueue()
    service = IngestionService(ingestion_engine, storage, queue)
    app.state.ingestion_service = service
    context = AuthContext(
        session_id=new_id(),
        workspace_id=workspace_id,
        user_id=user_id,
        display_name="Synthetic Ingestion Manager",
        role="manager",
        csrf_token="synthetic-csrf",
    )
    app.state.auth_service.authenticate = lambda *args, **kwargs: AuthenticationCheck(context, True)

    with TestClient(app, base_url="https://testserver") as client:
        app.state.file_storage = storage
        client.cookies.set(COOKIE_NAME, "synthetic-session", path="/api/v1")
        batch_response = client.post(
            "/api/v1/batches",
            json={"label": "Synthetic PO and GRN corrections"},
            headers={"X-CSRF-Token": "synthetic-csrf"},
        )
        assert batch_response.status_code == 201, batch_response.text
        batch_id = UUID(batch_response.json()["id"])

        def upload_and_parse(
            zone: str,
            name: str,
            content: bytes,
            job_base: int,
            mapping: dict[str, str] | None = None,
        ) -> UUID:
            uploaded = client.post(
                f"/api/v1/files?batch_id={batch_id}&zone={zone}",
                content=content,
                headers={
                    "Content-Type": "application/octet-stream",
                    "X-File-Name": name,
                    "X-CSRF-Token": "synthetic-csrf",
                },
            )
            assert uploaded.status_code == 201, uploaded.text
            file_id = UUID(uploaded.json()["file_id"])
            processed = client.post(
                f"/api/v1/batches/{batch_id}/process",
                headers={"X-CSRF-Token": "synthetic-csrf"},
            )
            assert processed.status_code == 202, processed.text
            service.ingest_file(file_id, job_id=job_base)
            service.parse_file(file_id, job_id=job_base + 1)
            if mapping is not None:
                saved = client.post(
                    f"/api/v1/files/{file_id}/column-mapping",
                    json={"column_mapping": mapping},
                    headers={"X-CSRF-Token": "synthetic-csrf"},
                )
                assert saved.status_code == 202, saved.text
                service.ingest_file(file_id, job_id=job_base + 2)
                service.parse_file(file_id, job_id=job_base + 3)
            return file_id

        po_file_id = upload_and_parse("po", "synthetic-po.csv", _PO_CSV, 1)
        grn_content = _BETA_GRN_CSV.replace(b"GRN-2026-777", b"GRN-2026-001").replace(
            b"PO-2026-777", b"PO-2026-001"
        )
        grn_file_id = upload_and_parse(
            "grn",
            "synthetic-grn.csv",
            grn_content,
            5,
            mapping={"PO Ref": "po_refs", "GRN No": "doc_no", "GRN Date": "doc_date"},
        )
        vendor_resolution = VendorResolutionService(ingestion_engine)
        vendor_resolution.resolve_workspace(workspace_id)
        mapping_service = MappingService(ingestion_engine, queue)
        mapping_service.resolve_workspace(workspace_id)

        documents = client.get(f"/api/v1/documents?batch_id={batch_id}")
        assert documents.status_code == 200, documents.text
        document_ids = {item["doc_type"]: UUID(item["id"]) for item in documents.json()}
        po_id = document_ids["po"]
        grn_id = document_ids["grn"]
        with ingestion_engine.connect() as connection:
            active_grn_po = connection.scalar(
                sa.select(metadata.tables["link_current"].c.current_link_id).where(
                    metadata.tables["link_current"].c.workspace_id == workspace_id,
                    metadata.tables["link_current"].c.from_document_id == grn_id,
                    metadata.tables["link_current"].c.to_document_id == po_id,
                    metadata.tables["link_current"].c.link_type == "grn_po",
                )
            )
            dirty_scope = metadata.tables["dirty_scope"]
            scope_before_po = connection.scalar(
                sa.select(dirty_scope.c.mark_count).where(
                    dirty_scope.c.workspace_id == workspace_id,
                    dirty_scope.c.scope_kind == "po",
                    dirty_scope.c.scope_id == po_id,
                    dirty_scope.c.processed_at.is_(None),
                )
            )
        assert active_grn_po is not None

        po_before = client.get(f"/api/v1/documents/{po_id}")
        assert po_before.status_code == 200, po_before.text
        po_old_version = UUID(po_before.json()["version"]["id"])
        po_line = po_before.json()["lines"][0]
        po_jobs_before = queue.jobs.count(("resolve_vendors", po_file_id))
        corrected_po = client.put(
            f"/api/v1/documents/{po_id}/correction",
            json={
                "revision": "Revision 2",
                "status": "closed",
                "lines": [{"id": po_line["id"], "line_no": po_line["line_no"], "quantity": "3"}],
                "verified": True,
                "reason": "Synthetic PO correction.",
            },
            headers={
                "If-Match": po_before.headers["etag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert corrected_po.status_code == 200, corrected_po.text
        po_saved = corrected_po.json()
        assert UUID(po_saved["version"]["id"]) != po_old_version
        assert po_saved["version"]["revision"] == "Revision 2"
        assert po_saved["version"]["po_status"] == "closed"
        assert po_saved["lines"][0]["quantity"] == "3.0000"
        assert po_saved["version"]["needs_verification"] is False
        assert {
            item["field_name"] for item in po_saved["field_evidence"] if item["source"] == "manual"
        } >= {
            "revision",
            "status",
            "quantity",
        }
        assert queue.jobs.count(("resolve_vendors", po_file_id)) == po_jobs_before + 1
        with ingestion_engine.connect() as connection:
            scope_after_po = connection.scalar(
                sa.select(dirty_scope.c.mark_count).where(
                    dirty_scope.c.workspace_id == workspace_id,
                    dirty_scope.c.scope_kind == "po",
                    dirty_scope.c.scope_id == po_id,
                    dirty_scope.c.processed_at.is_(None),
                )
            )
        assert scope_after_po is not None and scope_after_po > (scope_before_po or 0)

        grn_before = client.get(f"/api/v1/documents/{grn_id}")
        assert grn_before.status_code == 200, grn_before.text
        grn_old_version = UUID(grn_before.json()["version"]["id"])
        grn_line = grn_before.json()["lines"][0]
        grn_jobs_before = queue.jobs.count(("resolve_vendors", grn_file_id))
        scope_before_grn = scope_after_po
        corrected_grn = client.put(
            f"/api/v1/documents/{grn_id}/correction",
            json={
                "challan_no": "SYN-CHALLAN-001",
                "lines": [
                    {
                        "id": grn_line["id"],
                        "line_no": grn_line["line_no"],
                        "accepted_qty": "1",
                    }
                ],
                "verified": True,
                "reason": "Synthetic GRN correction.",
            },
            headers={
                "If-Match": grn_before.headers["etag"],
                "X-CSRF-Token": "synthetic-csrf",
            },
        )
        assert corrected_grn.status_code == 200, corrected_grn.text
        grn_saved = corrected_grn.json()
        assert UUID(grn_saved["version"]["id"]) != grn_old_version
        assert grn_saved["version"]["challan_no"] == "SYN-CHALLAN-001"
        assert grn_saved["lines"][0]["accepted_qty"] == "1.0000"
        assert grn_saved["version"]["needs_verification"] is False
        assert {
            item["field_name"] for item in grn_saved["field_evidence"] if item["source"] == "manual"
        } >= {
            "challan_no",
            "accepted_qty",
        }
        assert queue.jobs.count(("resolve_vendors", grn_file_id)) == grn_jobs_before + 1
        with ingestion_engine.connect() as connection:
            scope_after_grn = connection.scalar(
                sa.select(dirty_scope.c.mark_count).where(
                    dirty_scope.c.workspace_id == workspace_id,
                    dirty_scope.c.scope_kind == "po",
                    dirty_scope.c.scope_id == po_id,
                    dirty_scope.c.processed_at.is_(None),
                )
            )
        assert scope_after_grn is not None and scope_after_grn > (scope_before_grn or 0)


def test_rescan_fingerprint_creates_version_and_lower_confidence_stays_alternative(
    ingestion_engine: Engine,
    tmp_path: Any,
) -> None:
    workspace_id, user_id = _seed_workspace(ingestion_engine)
    service = IngestionService(
        ingestion_engine, LocalObjectStore(tmp_path / "objects"), _FakeQueue()
    )
    source = metadata.tables["source_file"]
    document_table = metadata.tables["document"]
    version_table = metadata.tables["document_version"]
    identity_table = metadata.tables["document_identity"]
    page_hashes = (PageImageHash(page=1, value="0123456789abcdef"),)

    def _value(value: Any, *, manual: bool) -> FieldValue[Any]:
        return FieldValue(
            value=value,
            confidence=Decimal("1.00") if manual else Decimal("0.80"),
            source=FieldSource.MANUAL if manual else FieldSource.RULE,
            evidence=FieldEvidence(page=1, snippet="synthetic identity value"),
            origin=FieldOrigin.MANUAL if manual else FieldOrigin.OCR,
            verification=FieldVerification.USER if manual else FieldVerification.NONE,
            verified_by="synthetic-owner" if manual else None,
            verified_at=datetime(2026, 4, 1, tzinfo=UTC) if manual else None,
        )

    def _invoice(
        *,
        manual: bool,
        grand_total: Decimal = Decimal("118.00"),
    ) -> Document:
        return Document(
            header=InvoiceHeader(
                doc_type=DocType.INVOICE,
                doc_no=_value("INV-SYN-AC8B-001", manual=manual),
                doc_date=_value(date(2026, 4, 1), manual=manual),
                vendor_name=_value("Synthetic Identity Supplier", manual=manual),
                grand_total=_value(grand_total, manual=manual),
            ),
            lines=(
                DocLine(
                    line_no=1,
                    description=_value("Synthetic identity item", manual=manual),
                    quantity=_value(Decimal("1"), manual=manual),
                    rate=_value(Decimal("100.00"), manual=manual),
                    amount=_value(Decimal("100.00"), manual=manual),
                ),
            ),
            source_locator="pdf:synthetic-ac8b:pages:1",
            needs_verification=not manual,
        )

    source_files: dict[int, tuple[UUID, str]] = {}

    def _source_file(index: int) -> tuple[UUID, str]:
        if index in source_files:
            return source_files[index]
        file_id = new_id()
        file_sha = hashlib.sha256(f"synthetic-ac8b-file-{index}".encode()).hexdigest()
        with ingestion_engine.begin() as connection:
            connection.execute(
                sa.insert(source).values(
                    id=file_id,
                    workspace_id=workspace_id,
                    sha256=file_sha,
                    display_name=f"synthetic-{index}.pdf",
                    zone="invoice",
                    zone_uploaded="invoice",
                    format="pdf",
                    status="parsed",
                    size_bytes=128,
                    object_key=f"synthetic/{file_id}",
                    uploaded_by=user_id,
                )
            )
        source_files[index] = (file_id, file_sha)
        return file_id, file_sha

    def _upsert(
        index: int,
        *,
        manual: bool,
        page_hashes_override: tuple[PageImageHash, ...] | None = None,
        grand_total: Decimal = Decimal("118.00"),
    ) -> tuple[str, UUID, UUID | None]:
        file_id, file_sha = _source_file(index)
        with ingestion_engine.begin() as connection:
            return service._upsert_document(
                connection,
                workspace_id=workspace_id,
                batch_id=None,
                file_id=file_id,
                actor_id=user_id,
                document=_invoice(
                    manual=manual,
                    grand_total=grand_total,
                ),
                file_sha=file_sha,
                page_hashes=(page_hashes if page_hashes_override is None else page_hashes_override),
                job_id=None,
                parser_version=PARSER_VERSION,
                layout_signature="a" * 64,
                extraction_method=(ExtractionMethod.MANUAL if manual else ExtractionMethod.PDF_OCR),
                uom_aliases={},
                charge_keywords=(),
            )

    first = _upsert(1, manual=False)
    second = _upsert(
        2,
        manual=False,
        page_hashes_override=(PageImageHash(page=1, value="0123456789abcdee"),),
    )
    verified = _upsert(3, manual=True)
    alternative = _upsert(4, manual=False)
    verified_changed_content = _upsert(8, manual=True, grand_total=Decimal("119.00"))
    different_page = _upsert(
        5,
        manual=False,
        page_hashes_override=(PageImageHash(page=1, value="fedcba9876543210"),),
    )
    different_content = _upsert(
        6,
        manual=False,
        page_hashes_override=(PageImageHash(page=1, value="fedcba9876543210"),),
        grand_total=Decimal("119.00"),
    )
    fallback_unhashed = _upsert(7, manual=False, page_hashes_override=())
    fallback_hashed = _upsert(
        7,
        manual=False,
        page_hashes_override=(PageImageHash(page=1, value="aaaaaaaaaaaaaaaa"),),
    )
    fallback_po = Document(
        header=POHeader(),
        source_locator="pdf:synthetic-ac8b:pages:1",
    )
    file_id, file_sha = _source_file(7)
    with ingestion_engine.begin() as connection:
        different_type_fallback = service._upsert_document(
            connection,
            workspace_id=workspace_id,
            batch_id=None,
            file_id=file_id,
            actor_id=user_id,
            document=fallback_po,
            file_sha=file_sha,
            page_hashes=(),
            job_id=None,
            parser_version=PARSER_VERSION,
            layout_signature="a" * 64,
            extraction_method=ExtractionMethod.PDF_OCR,
            uom_aliases={},
            charge_keywords=(),
        )

    assert first[0] == "new"
    assert second[0] == "updated" and second[1] == first[1]
    assert different_page[0] == "new" and different_page[1] != first[1]
    assert different_content[0] == "new" and different_content[1] != first[1]
    assert verified_changed_content[0] == "updated" and verified_changed_content[1] == first[1]
    assert verified[0] == "updated" and verified[1] == first[1]
    assert alternative[0] == "alternative" and alternative[1] == first[1]
    assert fallback_unhashed[0] == "new"
    assert fallback_hashed[0] == "skipped" and fallback_hashed[1] == fallback_unhashed[1]
    assert different_type_fallback[0] == "new"
    assert different_type_fallback[1] != fallback_unhashed[1]
    with ingestion_engine.connect() as connection:
        current_version_id = connection.scalar(
            sa.select(document_table.c.current_version_id).where(
                document_table.c.workspace_id == workspace_id,
                document_table.c.id == first[1],
            )
        )
        versions = connection.scalar(
            sa.select(sa.func.count())
            .select_from(version_table)
            .where(
                version_table.c.workspace_id == workspace_id,
                version_table.c.document_id == first[1],
            )
        )
        identities = connection.scalar(
            sa.select(sa.func.count())
            .select_from(identity_table)
            .join(version_table, version_table.c.id == identity_table.c.document_version_id)
            .where(
                identity_table.c.workspace_id == workspace_id,
                version_table.c.document_id == first[1],
            )
        )
    assert current_version_id == verified_changed_content[2]
    assert versions == 5
    assert identities == 5


def test_content_sniffing_and_zip_path_guards_use_bytes_and_archive_names() -> None:
    assert sniff_format(b"%PDF-1.7\n").value == "pdf"
    assert sniff_format(b"PK\x03\x04\x00").value == "zip"
    assert sniff_format(b"\xef\xbb\xbf<root/>").value == "xml"
    assert sniff_format("<root/>".encode("utf-16")).value == "xml"
    assert sniff_format(b"\xfe\xff" + "<root/>".encode("utf-16-be")).value == "xml"
    assert sniff_format("<root/>".encode("utf-16-le")).value == "xml"
    assert sniff_format(b"Invoice No,Date\nA-1,2026-04-01\n").value == "csv"
    assert sniff_document_type(_CSV, FileFormat.CSV).value == "invoice"
