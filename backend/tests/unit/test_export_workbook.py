"""Read-back checks for the P0 Excel workbook contract."""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import reconbench_app.export_service as export_service
from openpyxl import load_workbook
from reconbench_app.export_service import P0_SHEETS, ExportService


def _case(*, vendor_id: UUID, invoice_no: str = "0142") -> dict[str, Any]:
    return {
        "id": uuid4(),
        "effective_vendor_id": vendor_id,
        "link_summary": [],
        "invoice_taxable": Decimal("200.00"),
        "cgst_total": Decimal("10.00"),
        "sgst_total": Decimal("0"),
        "igst_total": Decimal("0"),
        "round_off": Decimal("0"),
        "invoice_total": Decimal("210.00"),
        "expected_taxable": Decimal("200.00"),
        "variance": Decimal("0"),
        "variance_pct": Decimal("0"),
        "payable_taxable": Decimal("180.00"),
        "at_risk": Decimal("20.00"),
        "effective_status": "exception",
        "status": "exception",
        "reason_codes": ["RATE_MISMATCH"],
        "payment_due_date": date(2026, 10, 15),
        "vendor_name": "=2+3",
        "effective_vendor_name": "=2+3",
        "vendor_gstin": "29ABCDE1234F1Z5",
        "invoice_no": invoice_no,
        "invoice_date": date(2026, 10, 1),
        "po_nos": ["0007"],
        "grn_nos": ["G-1"],
        "human_touched": False,
        "reading_method": "rules",
    }


def _context(vendor_ids: list[UUID]) -> dict[str, Any]:
    return {
        "workspace_id": uuid4(),
        "run_id": uuid4(),
        "filters": {"from": "2026-10-01", "to": "2026-10-31"},
        "run": {
            "finished_at": datetime(2026, 10, 2),
            "engine_version": "0.1",
            "rules_hash": "synthetic-rules-hash",
            "config_version_id": str(uuid4()),
            "ai_mode": "rules_only",
            "ai_provider": None,
            "ai_model": None,
            "config_body": {"engine": {"tolerances": {"qty": {"pct": 1}}}},
        },
        "vendor_doc_counts": {
            vendor_id: Counter({"po": 1, "grn": 1, "invoice": 1}) for vendor_id in vendor_ids
        },
        "gstins": {vendor_id: ["29ABCDE1234F1Z5"] for vendor_id in vendor_ids},
    }


def _write(path: Path, cases: list[dict[str, Any]], vendor_ids: list[UUID]) -> None:
    ExportService._write_workbook(
        ExportService.__new__(ExportService),
        path,
        cases=cases,
        lines=[],
        context=_context(vendor_ids),
    )


def test_workbook_readback_checks_formulas_cached_values_and_text_safety(tmp_path: Path) -> None:
    vendor_id = uuid4()
    path = tmp_path / "reconciliation.xlsx"
    _write(path, [_case(vendor_id=vendor_id)], [vendor_id])

    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    try:
        assert tuple(formulas.sheetnames) == P0_SHEETS
        assert formulas["Summary"]["A1"].value == "Metric"
        assert formulas["Invoice Recon"]["C1"].value == "Invoice No."
        assert formulas["Invoice Recon"]["C2"].value == "0142"
        assert formulas["Invoice Recon"]["A2"].value == "=2+3"
        assert formulas["Invoice Recon"]["A2"].data_type == "s"
        assert formulas["Invoice Recon"]["K2"].value == "=I2+J2"
        assert formulas["Invoice Recon"]["M2"].value == "=I2-L2"
        assert formulas["Invoice Recon"]["P2"].value.startswith("=IF(OR(")
        assert formulas["Invoice Recon"].freeze_panes == "A2"
        assert values["Invoice Recon"]["K2"].value == 210
        assert values["Invoice Recon"]["P2"].value == 20
        assert values["Summary"]["B11"].value == 210
        assert values["Summary"]["B12"].value == 180
        assert values["Summary"]["B13"].value == 20
    finally:
        formulas.close()
        values.close()


def test_workbook_splits_rows_and_summary_formulas_cover_each_invoice_sheet(
    tmp_path: Path, monkeypatch: Any
) -> None:
    monkeypatch.setattr(export_service, "EXCEL_MAX_ROWS", 2)
    vendor_ids = [uuid4(), uuid4(), uuid4()]
    cases = [
        _case(vendor_id=vendor_id, invoice_no=f"INV-{index}")
        for index, vendor_id in enumerate(vendor_ids)
    ]
    path = tmp_path / "split.xlsx"
    _write(path, cases, vendor_ids)

    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    try:
        for sheet_name in P0_SHEETS:
            assert sheet_name in formulas.sheetnames
        assert [name for name in formulas.sheetnames if name.startswith("Invoice Recon")] == [
            "Invoice Recon",
            "Invoice Recon 2",
            "Invoice Recon 3",
        ]
        assert [name for name in formulas.sheetnames if name.startswith("Exceptions")] == [
            "Exceptions",
            "Exceptions 2",
            "Exceptions 3",
        ]
        assert [name for name in formulas.sheetnames if name.startswith("Vendor Summary")] == [
            "Vendor Summary",
            "Vendor Summary 2",
            "Vendor Summary 3",
        ]
        summary_formula = formulas["Summary"]["B11"].value
        assert "'Invoice Recon'!K2:K2" in summary_formula
        assert "'Invoice Recon 2'!K2:K2" in summary_formula
        assert "'Invoice Recon 3'!K2:K2" in summary_formula
        assert values["Summary"]["B11"].value == 630
        assert values["Summary"]["B12"].value == 540
        assert values["Summary"]["B13"].value == 60
    finally:
        formulas.close()
        values.close()


def test_summary_counts_needs_review_apart_from_exceptions(tmp_path: Path) -> None:
    vendor_id = uuid4()
    needs_review = _case(vendor_id=vendor_id, invoice_no="SCAN-1") | {
        "effective_status": "needs_review",
        "status": "needs_review",
        "reason_codes": ["NO_TEXT_LAYER"],
    }
    path = tmp_path / "statuses.xlsx"
    _write(path, [_case(vendor_id=vendor_id), needs_review], [vendor_id])

    values = load_workbook(path, data_only=True)
    try:
        summary = {
            row[0]: row[1]
            for row in values["Summary"].iter_rows(min_row=2, max_col=2, values_only=True)
        }
        assert summary["Invoices"] == 2
        assert summary["Exceptions"] == 1
        assert summary["Needs Review"] == 1
        assert summary["Duplicate"] == 0
    finally:
        values.close()


def test_invoice_sheet_names_the_reading_method_in_the_last_column(tmp_path: Path) -> None:
    """PD-2c / Q13: each invoice row says how its version was read; no column moves."""

    vendor_id = uuid4()
    path = tmp_path / "reconciliation.xlsx"
    ai_case = {**_case(vendor_id=vendor_id, invoice_no="0143"), "reading_method": "ai"}
    _write(path, [_case(vendor_id=vendor_id), ai_case], [vendor_id])

    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Invoice Recon"]
        assert sheet["T1"].value == "Payment Due Date"
        assert sheet["U1"].value == "Reading Method"
        assert sheet.max_column == 21
        assert [sheet["C2"].value, sheet["U2"].value] == ["0142", "Rules"]
        assert [sheet["C3"].value, sheet["U3"].value] == ["0143", "AI reading"]
        assert sheet["K2"].value == "=I2+J2"
    finally:
        workbook.close()


def test_field_origin_sheet_reports_origins_and_verification_with_text_safety(
    tmp_path: Path,
) -> None:
    """PD-4c / D7: per-field origin sheet reports rules, AI, person, and import origins.

    Also tests text safety (formula injection prevention).
    """
    vendor_id = uuid4()
    version_id = uuid4()
    path = tmp_path / "origins.xlsx"
    case = _case(vendor_id=vendor_id, invoice_no="=INV-001") | {
        "invoice_version_id": version_id,
        "reading_method": "rules",
    }
    field_evidence = {
        version_id: {
            "doc_no": {"origin": "ocr", "verification": "none"},
            "doc_date": {"origin": "ai", "verification": "none"},
            "grand_total": {"origin": "manual", "verification": "user"},
            "vendor_name": {"origin": "erp_import", "verification": "system_check"},
        }
    }
    context = _context([vendor_id]) | {"field_evidence": field_evidence}
    ExportService._write_workbook(
        ExportService.__new__(ExportService),
        path,
        cases=[case],
        lines=[],
        context=context,
    )

    formulas = load_workbook(path, data_only=False)
    try:
        assert "Field Origin" in formulas.sheetnames
        sheet = formulas["Field Origin"]
        headers = [sheet.cell(1, col).value for col in range(1, 5)]
        assert headers == ["Invoice No.", "Field", "Origin", "Verified"]
        rows = {
            sheet.cell(row, 2).value: (
                sheet.cell(row, 1).value,
                sheet.cell(row, 3).value,
                sheet.cell(row, 4).value,
            )
            for row in range(2, sheet.max_row + 1)
        }
        # Rules field
        assert rows["doc_no"] == ("=INV-001", "Rules", "No")
        # Unverified AI field
        assert rows["doc_date"] == ("=INV-001", "AI", "No")
        # Person-corrected field
        assert rows["grand_total"] == ("=INV-001", "Person", "Yes")
        # Import field
        assert rows["vendor_name"] == ("=INV-001", "Import", "Yes")
        # Text safety on invoice_no leading '=': string cell, not formula
        assert sheet["A2"].value == "=INV-001"
        assert sheet["A2"].data_type == "s"
    finally:
        formulas.close()
