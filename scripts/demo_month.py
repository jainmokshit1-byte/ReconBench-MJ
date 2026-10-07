"""Run the synthetic demo month end to end through a disposable ReconBench stack.

Usage:
    python scripts/demo_month.py render --out DIR   # write the month's upload files
    python scripts/demo_month.py run                # full stack smoke (``make demo``)

``run`` builds the backend image from the working tree, starts throwaway
PostgreSQL, MinIO, API, and worker containers on a private Docker network,
uploads the month through the HTTP API exactly as a user would, reconciles it,
exports the workbook, and compares every case with ``expected.yaml``. All
containers and the network are removed afterwards. Nothing here touches the
local ``make up`` stack, and no external provider is called.
"""

from __future__ import annotations

import argparse
import io
import json
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import Counter
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import httpx
import yaml

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "fixtures" / "demo-month"
CENT = Decimal("0.01")
POSTGRES_IMAGE = "postgres:16.15-bookworm"
DEMO_IMAGE = "reconbench-demo:local"
MINIO_IMAGE = "reconbench-minio:latest"
TERMINAL_FILE_STATES = {"parsed", "partial", "failed", "needs_password", "removed"}


# --------------------------------------------------------------------------- spec


@dataclass(frozen=True)
class Line:
    code: str
    description: str
    hsn: str
    uom: str
    qty: Decimal
    rate: Decimal
    gst: Decimal

    @property
    def amount(self) -> Decimal:
        return (self.qty * self.rate).quantize(CENT)


@dataclass(frozen=True)
class Charge:
    description: str
    hsn: str
    amount: Decimal
    gst: Decimal


@dataclass(frozen=True)
class Totals:
    taxable: Decimal
    cgst: Decimal
    sgst: Decimal
    igst: Decimal

    @property
    def grand(self) -> Decimal:
        return self.taxable + self.cgst + self.sgst + self.igst


class Month:
    def __init__(self, raw: dict[str, Any]) -> None:
        self.raw = raw
        self.buyer = raw["buyer"]
        self.vendors = raw["vendors"]
        self.pos = {item["number"]: item for item in raw["purchase_orders"]}
        self.grns = {item["number"]: item for item in raw["grns"]}
        self.invoices = {item["number"]: item for item in raw["invoices"]}
        self.catalog: dict[str, dict[str, Any]] = {}
        for po in raw["purchase_orders"]:
            for line in po["lines"]:
                self.catalog[line["code"]] = line

    def vendor(self, document: dict[str, Any]) -> dict[str, str]:
        return self.vendors[document["vendor"]]

    def lines(self, document: dict[str, Any]) -> list[Line]:
        result = []
        for line in document["lines"]:
            item = self.catalog[line["code"]]
            result.append(
                Line(
                    code=line["code"],
                    description=item["description"],
                    hsn=item["hsn"],
                    uom=item["uom"],
                    qty=Decimal(line.get("qty", item["qty"])),
                    rate=Decimal(line.get("rate", item["rate"])),
                    gst=Decimal(item["gst"]),
                )
            )
        return result

    def charges(self, document: dict[str, Any]) -> list[Charge]:
        return [
            Charge(
                description=item["description"],
                hsn=item["hsn"],
                amount=Decimal(item["amount"]),
                gst=Decimal(item["gst"]),
            )
            for item in document.get("charges", [])
        ]

    def totals(self, document: dict[str, Any]) -> Totals:
        lines = self.lines(document)
        charges = self.charges(document)
        taxable = sum((line.amount for line in lines), Decimal("0")) + sum(
            (charge.amount for charge in charges), Decimal("0")
        )
        tax = sum((line.amount * line.gst / 100 for line in lines), Decimal("0")) + sum(
            (charge.amount * charge.gst / 100 for charge in charges), Decimal("0")
        )
        tax = tax.quantize(CENT, ROUND_HALF_UP)
        if self.intrastate(document):
            half = (tax / 2).quantize(CENT, ROUND_HALF_UP)
            return Totals(taxable.quantize(CENT), half, tax - half, Decimal("0.00"))
        return Totals(taxable.quantize(CENT), Decimal("0.00"), Decimal("0.00"), tax)

    def intrastate(self, document: dict[str, Any]) -> bool:
        return bool(self.vendor(document)["gstin"][:2] == self.buyer["gstin"][:2])


def load_month() -> Month:
    return Month(yaml.safe_load((FIXTURE / "month.yaml").read_text(encoding="utf-8")))


def _dmy(value: str) -> str:
    year, month, day = value.split("-")
    return f"{day}/{month}/{year}"


def _money(value: Decimal) -> str:
    return f"{value:.2f}"


def _qty(value: Decimal) -> str:
    return f"{value.normalize():f}"


# ------------------------------------------------------------------- renderers


def render_tally_xml(month: Month, numbers: list[str], kind: str) -> bytes:
    vouchers = []
    for number in numbers:
        document = month.pos[number] if kind == "po" else month.invoices[number]
        vendor = month.vendor(document)
        totals = month.totals(document)
        entries = []
        for line in month.lines(document):
            entries.append(
                "<ALLINVENTORYENTRIES.LIST>"
                f"<STOCKITEMNAME>{escape(line.description)}</STOCKITEMNAME>"
                f"<STOCKITEMCODE>{escape(line.code)}</STOCKITEMCODE>"
                f"<HSNCODE>{line.hsn}</HSNCODE>"
                f"<ACTUALQTY>{_qty(line.qty)} {line.uom}</ACTUALQTY>"
                f"<RATE>{_money(line.rate)}/{line.uom}</RATE>"
                f"<AMOUNT>{_money(line.amount)}</AMOUNT>"
                f"<GSTRATE>{_qty(line.gst)}</GSTRATE>"
                "</ALLINVENTORYENTRIES.LIST>"
            )
        for charge in month.charges(document):
            entries.append(
                "<ALLINVENTORYENTRIES.LIST>"
                f"<STOCKITEMNAME>{escape(charge.description)}</STOCKITEMNAME>"
                f"<HSNCODE>{charge.hsn}</HSNCODE>"
                f"<AMOUNT>{_money(charge.amount)}</AMOUNT>"
                f"<GSTRATE>{_qty(charge.gst)}</GSTRATE>"
                "</ALLINVENTORYENTRIES.LIST>"
            )
        voucher_type = "Purchase Order" if kind == "po" else "Purchase"
        reference = "" if kind == "po" else f"<POREFERENCE>{document['po']}</POREFERENCE>"
        vouchers.append(
            "<VOUCHER>"
            f"<VOUCHERTYPENAME>{voucher_type}</VOUCHERTYPENAME>"
            f"<VOUCHERNUMBER>{escape(number)}</VOUCHERNUMBER>"
            f"<DATE>{document['date'].replace('-', '')}</DATE>"
            f"<PARTYLEDGERNAME>{escape(vendor['name'])}</PARTYLEDGERNAME>"
            f"<PARTYGSTIN>{vendor['gstin']}</PARTYGSTIN>"
            f"<BUYERGSTIN>{month.buyer['gstin']}</BUYERGSTIN>"
            f"{reference}"
            f"<TAXABLETOTAL>{_money(totals.taxable)}</TAXABLETOTAL>"
            f"<CGST>{_money(totals.cgst)}</CGST><SGST>{_money(totals.sgst)}</SGST>"
            f"<IGST>{_money(totals.igst)}</IGST>"
            f"<VOUCHERTOTAL>{_money(totals.grand)}</VOUCHERTOTAL>"
            f"{''.join(entries)}"
            "</VOUCHER>"
        )
    body = "\n".join(vouchers)
    return f"<ENVELOPE><BODY>\n{body}\n</BODY></ENVELOPE>\n".encode()


def _csv(rows: list[list[str]]) -> bytes:
    import csv

    stream = io.StringIO()
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerows(rows)
    return stream.getvalue().encode()


def render_invoice_csv(month: Month, numbers: list[str]) -> bytes:
    rows = [
        [
            "Invoice No", "Invoice Date", "Purchase Order", "Vendor Name", "Vendor GSTIN",
            "Buyer GSTIN", "Currency", "Taxable Total", "CGST", "SGST", "IGST",
            "Grand Total", "Line No", "Item Code", "Description", "HSN", "UOM",
            "Quantity", "Rate", "Amount", "GST Rate",
        ]
    ]  # fmt: skip
    for number in numbers:
        document = month.invoices[number]
        vendor = month.vendor(document)
        totals = month.totals(document)
        header = [
            number, _dmy(document["date"]), document["po"], vendor["name"], vendor["gstin"],
            month.buyer["gstin"], "INR", _money(totals.taxable), _money(totals.cgst),
            _money(totals.sgst), _money(totals.igst), _money(totals.grand),
        ]  # fmt: skip
        index = 0
        for index, line in enumerate(month.lines(document), start=1):
            rows.append(
                [
                    *header,
                    str(index),
                    line.code,
                    line.description,
                    line.hsn,
                    line.uom,
                    _qty(line.qty),
                    _money(line.rate),
                    _money(line.amount),
                    _qty(line.gst),
                ]  # fmt: skip
            )
        for offset, charge in enumerate(month.charges(document), start=index + 1):
            rows.append(
                [
                    *header,
                    str(offset),
                    "",
                    charge.description,
                    charge.hsn,
                    "",
                    "",
                    "",
                    _money(charge.amount),
                    _qty(charge.gst),
                ]  # fmt: skip
            )
    return _csv(rows)


def render_grn_csv(month: Month, numbers: list[str]) -> bytes:
    rows = [
        [
            "GRN No", "GRN Date", "Purchase Order", "Vendor Name", "Vendor GSTIN", "Line No",
            "Item Code", "Description", "HSN", "UOM", "Received Qty", "Accepted Qty",
            "Rejected Qty",
        ]
    ]  # fmt: skip
    for number in numbers:
        grn = month.grns[number]
        po = month.pos[grn["po"]]
        vendor = month.vendor(po)
        for index, line in enumerate(grn["lines"], start=1):
            item = month.catalog[line["code"]]
            rows.append(
                [
                    number,
                    _dmy(grn["date"]),
                    grn["po"],
                    vendor["name"],
                    vendor["gstin"],
                    str(index),
                    line["code"],
                    item["description"],
                    item["hsn"],
                    item["uom"],
                    line["received"],
                    line["accepted"],
                    line.get("rejected", "0"),
                ]  # fmt: skip
            )
    return _csv(rows)


_ONES = (
    "", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten",
    "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen",
    "Eighteen", "Nineteen",
)  # fmt: skip
_TENS = (
    "",
    "",
    "Twenty",
    "Thirty",
    "Forty",
    "Fifty",
    "Sixty",
    "Seventy",
    "Eighty",
    "Ninety",
)


def _below_thousand(value: int) -> list[str]:
    words: list[str] = []
    if value >= 100:
        words += [_ONES[value // 100], "Hundred"]
        value %= 100
    if value >= 20:
        words.append(_TENS[value // 10])
        value %= 10
    if value:
        words.append(_ONES[value])
    return words


def rupees_in_words(amount: Decimal) -> str:
    """Indian-system words, e.g. ``Rupees One Lakh Two Thousand Only``."""

    rupees = int(amount)
    paise = int((amount - rupees) * 100)
    words: list[str] = []
    for unit, size in (("Crore", 10_000_000), ("Lakh", 100_000), ("Thousand", 1_000)):
        if rupees >= size:
            words += [*_below_thousand(rupees // size), unit]
            rupees %= size
    words += _below_thousand(rupees)
    text = "Rupees " + " ".join(words)
    if paise:
        text += " and " + " ".join(_below_thousand(paise)) + " Paise"
    return text + " Only"


def render_invoice_pdf(month: Month, number: str) -> bytes:
    """A conventional Indian GST tax invoice with a real text layer."""

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle

    document = month.invoices[number]
    vendor = month.vendor(document)
    totals = month.totals(document)
    stream = io.BytesIO()
    page = canvas.Canvas(stream, pagesize=A4)
    width, height = A4
    page.setFont("Helvetica-Bold", 16)
    page.drawCentredString(width / 2, height - 50, "TAX INVOICE")
    page.setFont("Helvetica-Bold", 11)
    page.drawString(40, height - 85, vendor["name"])
    page.setFont("Helvetica", 9)
    page.drawString(40, height - 100, "Plot 14, Industrial Area Phase II, Pune 411026")
    page.drawString(40, height - 114, f"GSTIN: {vendor['gstin']}")
    right = 340
    page.setFont("Helvetica", 10)
    page.drawString(right, height - 85, f"Invoice No: {number}")
    page.drawString(right, height - 100, f"Invoice Date: {_dmy(document['date'])}")
    page.drawString(right, height - 115, f"PO No: {document['po']}")
    page.setFont("Helvetica", 9)
    page.drawString(40, height - 145, f"Bill To: {month.buyer['name']}")
    page.drawString(40, height - 159, f"Buyer GSTIN: {month.buyer['gstin']}")

    rows = [["Sl No", "Description", "HSN/SAC", "Qty", "UOM", "Rate", "Amount"]]
    index = 0
    for index, line in enumerate(month.lines(document), start=1):
        rows.append(
            [
                str(index),
                line.description,
                line.hsn,
                _qty(line.qty),
                line.uom,
                _money(line.rate),
                _money(line.amount),
            ]  # fmt: skip
        )
    for offset, charge in enumerate(month.charges(document), start=index + 1):
        rows.append(
            [
                str(offset),
                charge.description,
                charge.hsn,
                "",
                "",
                "",
                _money(charge.amount),
            ]
        )
    table = Table(rows, colWidths=[36, 200, 60, 44, 40, 64, 76])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
            ]
        )
    )
    _w, table_height = table.wrapOn(page, width - 80, 400)
    top = height - 185
    table.drawOn(page, 40, top - table_height)
    y = top - table_height - 24
    summary = [("Taxable Value", totals.taxable)]
    if month.intrastate(document):
        summary += [("CGST @ 9%", totals.cgst), ("SGST @ 9%", totals.sgst)]
    else:
        summary += [("IGST @ 18%", totals.igst)]
    summary += [("Grand Total", totals.grand)]
    for label, value in summary:
        page.setFont("Helvetica-Bold" if label == "Grand Total" else "Helvetica", 10)
        page.drawString(360, y, label)
        page.drawRightString(width - 40, y, _money(value))
        y -= 16
    page.setFont("Helvetica", 9)
    page.drawString(40, y - 10, f"Amount in words: {rupees_in_words(totals.grand)}")
    page.drawString(width - 220, y - 60, f"For {vendor['name']}")
    page.drawString(width - 220, y - 90, "Authorised Signatory")
    page.showPage()
    page.save()
    return stream.getvalue()


def _image_pdf(image: Any) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    page = canvas.Canvas(stream, pagesize=A4)
    width, height = A4
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=85)
    buffer.seek(0)
    page.drawImage(ImageReader(buffer), 0, 0, width=width, height=height)
    page.showPage()
    page.save()
    return stream.getvalue()


def render_scanned_pdf(month: Month, number: str) -> bytes:
    """Rasterise the invoice like an office scanner: 200 DPI grey, slight skew."""

    import pypdfium2  # type: ignore[import-untyped]
    from PIL import ImageFilter

    pdf = pypdfium2.PdfDocument(render_invoice_pdf(month, number))
    try:
        image = pdf[0].render(scale=200 / 72).to_pil().convert("L")
    finally:
        pdf.close()
    image = image.rotate(0.6, expand=False, fillcolor=255).filter(ImageFilter.GaussianBlur(0.5))
    return _image_pdf(image)


def render_unreadable_pdf() -> bytes:
    """A scan with no legible text: faint smudges on grey paper."""

    import random

    from PIL import Image, ImageDraw, ImageFilter

    generator = random.Random(19)
    image = Image.new("L", (1654, 2339), 235)
    draw = ImageDraw.Draw(image)
    for _ in range(60):
        x, y = generator.randrange(0, 1600), generator.randrange(0, 2300)
        draw.ellipse(
            (x, y, x + generator.randrange(20, 120), y + generator.randrange(5, 30)),
            fill=200,
        )
    return _image_pdf(image.filter(ImageFilter.GaussianBlur(6)))


def render_month(month: Month, out: Path) -> list[tuple[Path, str]]:
    out.mkdir(parents=True, exist_ok=True)
    uploads: list[tuple[Path, str]] = []
    for item in month.raw["files"]:
        content = _render_file(month, item)
        path = out / item["name"]
        path.write_bytes(content)
        uploads.append((path, item["zone"]))
    return uploads


def _render_file(month: Month, item: dict[str, Any]) -> bytes:
    render = item["render"]
    numbers = item.get("documents", [])
    if render == "tally_xml":
        return render_tally_xml(month, numbers, item["zone"])
    if render == "csv":
        if item["zone"] == "grn":
            return render_grn_csv(month, numbers)
        return render_invoice_csv(month, numbers)
    if render == "pdf_text":
        return render_invoice_pdf(month, numbers[0])
    if render == "pdf_scan":
        return render_scanned_pdf(month, numbers[0])
    if render == "pdf_unreadable":
        return render_unreadable_pdf()
    if render == "zip":
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for member in item["members"]:
                archive.writestr(
                    member["name"],
                    _render_file(month, {**member, "zone": item["zone"]}),
                )
        return stream.getvalue()
    raise ValueError(f"unknown render kind {render!r}")


# ------------------------------------------------------------------- the stack


def _docker(*args: str, check: bool = True, capture: bool = True) -> str:
    completed = subprocess.run(
        ["docker", *args],
        check=False,
        capture_output=capture,
        text=True,
        encoding="utf-8",
    )
    if check and completed.returncode != 0:
        raise RuntimeError(f"docker {args[0]} failed: {(completed.stderr or '').strip()[-2000:]}")
    return (completed.stdout or "").strip()


def _free_port() -> int:
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        return int(probe.getsockname()[1])


def _step(message: str) -> None:
    print(f"[demo] {message}", flush=True)


_SETUP = r"""
import os, sys, time, subprocess
import boto3, sqlalchemy as sa
from argon2 import PasswordHasher
url = os.environ["RB_DATABASE_URL"]
engine = sa.create_engine(url)
for _ in range(60):
    try:
        with engine.connect() as c:
            c.execute(sa.text("select 1"))
        break
    except Exception:
        time.sleep(1)
else:
    sys.exit("database never became ready")
run = lambda *a: subprocess.run(a, check=True)
run("/app/backend/.venv/bin/alembic", "-c", "/app/backend/alembic.ini", "upgrade", "head")
run("/app/backend/.venv/bin/procrastinate", "--app",
    "reconbench_app.db.procrastinate_schema.app", "schema", "--apply")
os.environ["RB_BOOTSTRAP_PASSWORD_HASH"] = PasswordHasher().hash(os.environ["DEMO_PASSWORD"])
from reconbench_app.db.seed import BootstrapSettings, seed_database
seed_database(engine, BootstrapSettings.from_environment())
s3 = boto3.client("s3", endpoint_url=os.environ["RB_S3_ENDPOINT_URL"],
    region_name=os.environ["RB_S3_REGION"],
    aws_access_key_id=os.environ["RB_S3_API_ACCESS_KEY"],
    aws_secret_access_key=os.environ["RB_S3_API_SECRET_KEY"])
for _ in range(60):
    try:
        s3.create_bucket(Bucket=os.environ["RB_S3_BUCKET"])
        break
    except Exception as exc:
        if "BucketAlreadyOwnedByYou" in str(exc):
            break
        time.sleep(1)
else:
    sys.exit("object store never became ready")
print("setup complete")
"""


class Stack:
    """Disposable containers on a private network; removed on exit."""

    def __init__(self, image: str) -> None:
        self.tag = f"rbdemo-{secrets.token_hex(4)}"
        self.image = image
        self.network = self.tag
        self.containers: list[str] = []
        self.db_password = secrets.token_urlsafe(18)
        self.s3_secret = secrets.token_urlsafe(18)
        self.email = "demo.manager@example.invalid"
        self.password = secrets.token_urlsafe(18)
        self.api_port = _free_port()

    @property
    def app_env(self) -> dict[str, str]:
        return {
            "RB_ENV": "local",
            "RB_DATABASE_URL": (
                f"postgresql+psycopg://demo:{self.db_password}@{self.tag}-postgres:5432/reconbench"
            ),
            "RB_S3_ENDPOINT_URL": f"http://{self.tag}-minio:9000",
            "RB_S3_REGION": "us-east-1",
            "RB_S3_BUCKET": "reconbench-demo",
            "RB_S3_API_ACCESS_KEY": "demo-minio",
            "RB_S3_API_SECRET_KEY": self.s3_secret,
            "RB_PUBLIC_BASE_URL": f"http://127.0.0.1:{self.api_port}",
            "RB_PDF_OCR_LANGUAGES": "eng",
            "RB_AI_EXTERNAL_CALLS": "disabled",
            "RB_LOG_LEVEL": "INFO",
        }

    def _env_args(self, extra: dict[str, str] | None = None) -> list[str]:
        args: list[str] = []
        for key, value in {**self.app_env, **(extra or {})}.items():
            args += ["-e", f"{key}={value}"]
        return args

    def _run(self, name: str, *args: str) -> None:
        container = f"{self.tag}-{name}"
        _docker("run", "-d", "--name", container, "--network", self.network, *args)
        self.containers.append(container)

    def start(self) -> None:
        _docker("network", "create", self.network)
        self._run(
            "postgres",
            "-e", "POSTGRES_DB=reconbench", "-e", "POSTGRES_USER=demo",
            "-e", f"POSTGRES_PASSWORD={self.db_password}",
            "--tmpfs", "/var/lib/postgresql/data",
            POSTGRES_IMAGE,
        )  # fmt: skip
        self._run(
            "minio",
            "-e", "MINIO_ROOT_USER=demo-minio", "-e", f"MINIO_ROOT_PASSWORD={self.s3_secret}",
            "--tmpfs", "/data",
            MINIO_IMAGE, "server", "/data",
        )  # fmt: skip
        _step("applying migrations, queue schema, seed, and bucket")
        setup = subprocess.run(
            [
                "docker", "run", "--rm", "--network", self.network,
                *self._env_args(
                    {
                        "DEMO_PASSWORD": self.password,
                        "RB_BOOTSTRAP_EMAIL": self.email,
                        "RB_BOOTSTRAP_DISPLAY_NAME": "Demo Manager",
                        "RB_BOOTSTRAP_WORKSPACE_NAME": "Demo Month",
                    }
                ),
                self.image, "/app/backend/.venv/bin/python", "-c", _SETUP,
            ],
            capture_output=True, text=True, encoding="utf-8", check=False,
        )  # fmt: skip
        if setup.returncode != 0:
            raise RuntimeError(
                f"stack setup failed:\n{setup.stdout[-2000:]}\n{setup.stderr[-3000:]}"
            )
        self._run(
            "api", "-p", f"127.0.0.1:{self.api_port}:8000",
            *self._env_args({"RB_LOG_LEVEL": "WARNING"}), self.image,
        )  # fmt: skip
        self._run(
            "worker", *self._env_args(), self.image,
            "/app/backend/.venv/bin/python", "-m", "reconbench_app.worker",
        )  # fmt: skip

    def queue_busy(self) -> int:
        output = _docker(
            "exec", f"{self.tag}-postgres", "psql", "-U", "demo", "-d", "reconbench", "-tAc",
            "SELECT count(*) FROM procrastinate_jobs WHERE status IN ('todo','doing')",
        )  # fmt: skip
        return int(output or "0")

    def logs(self, name: str) -> str:
        return _docker("logs", "--tail", "60", f"{self.tag}-{name}", check=False)

    def stop(self) -> None:
        for container in reversed(self.containers):
            _docker("rm", "-f", container, check=False)
        _docker("network", "rm", self.network, check=False)


# ------------------------------------------------------------------ the client


class Client:
    def __init__(self, base_url: str) -> None:
        self.http = httpx.Client(base_url=base_url, timeout=60.0)
        self.csrf = ""

    def wait_ready(self, seconds: float = 90.0) -> None:
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            try:
                if self.http.get("/healthz").status_code == 200:
                    return
            except httpx.HTTPError:
                pass
            time.sleep(1)
        raise RuntimeError("API did not become healthy")

    def login(self, email: str, password: str) -> None:
        response = self.http.post("/api/v1/auth/login", json={"email": email, "password": password})
        _expect(response, 200)
        self.csrf = response.json()["csrf_token"]

    def get(self, path: str, **kwargs: Any) -> httpx.Response:
        return self.http.get(path, **kwargs)

    def post(
        self, path: str, *, headers: dict[str, str] | None = None, **kwargs: Any
    ) -> httpx.Response:
        return self.http.post(
            path, headers={"X-CSRF-Token": self.csrf, **(headers or {})}, **kwargs
        )


def _expect(response: httpx.Response, *codes: int) -> None:
    if response.status_code not in codes:
        raise RuntimeError(
            f"{response.request.method} {response.request.url.path} -> {response.status_code}: "
            f"{response.text[:800]}"
        )


def _poll(description: str, check: Any, seconds: float, interval: float = 2.0) -> Any:
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        value = check()
        if value:
            return value
        time.sleep(interval)
    raise RuntimeError(f"timed out waiting for {description}")


# ------------------------------------------------------------------- the flow


def _items(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return payload
    for key in ("items", "results", "data"):
        if isinstance(payload.get(key), list):
            return list(payload[key])
    raise RuntimeError(f"unexpected list payload keys: {sorted(payload)}")


def _all(client: Client, path: str) -> list[dict[str, Any]]:
    """Every row of a list endpoint, following keyset cursors when present."""

    separator = "&" if "?" in path else "?"
    rows: list[dict[str, Any]] = []
    cursor: str | None = None
    while True:
        params = {"limit": 100, **({"cursor": cursor} if cursor else {})}
        response = client.get(
            f"{path}{separator}" + "&".join(f"{k}={v}" for k, v in params.items())
        )
        _expect(response, 200)
        payload = response.json()
        rows += _items(payload)
        cursor = payload.get("next_cursor") if isinstance(payload, dict) else None
        if not cursor:
            return rows


def drive(
    client: Client, stack: Stack, uploads: list[tuple[Path, str]], month: Month
) -> dict[str, Any]:
    batch = client.post("/api/v1/batches", json={"label": "Demo month September 2026"})
    _expect(batch, 201)
    batch_id = batch.json()["id"]
    for path, zone in uploads:
        response = client.post(
            f"/api/v1/files?batch_id={batch_id}&zone={zone}",
            content=path.read_bytes(),
            headers={
                "Content-Type": "application/octet-stream",
                "X-File-Name": path.name,
            },
        )
        _expect(response, 201)
    _step(f"uploaded {len(uploads)} files; processing")
    _expect(client.post(f"/api/v1/batches/{batch_id}/process"), 202)

    def files_done() -> dict[str, Any] | None:
        summary = client.get(f"/api/v1/batches/{batch_id}/summary").json()
        counts = summary.get("file_status_counts", {})
        pending = {
            state: count for state, count in counts.items() if state not in TERMINAL_FILE_STATES
        }
        return None if pending or not counts else summary

    try:
        summary = _poll("files to finish parsing", files_done, 600)
    except RuntimeError:
        print(json.dumps(client.get(f"/api/v1/batches/{batch_id}/summary").json(), indent=2))
        raise
    _poll("background jobs to drain", lambda: stack.queue_busy() == 0, 300)
    files = _all(client, f"/api/v1/files?batch_id={batch_id}")
    documents = _all(client, f"/api/v1/documents?batch_id={batch_id}")

    run = client.post(
        "/api/v1/runs",
        json={"run_date": month.raw["run_date"]},
        headers={"Idempotency-Key": f"demo-run-{stack.tag}"},
    )
    _expect(run, 202)
    run_id = run.json()["id"]
    _step("reconciling")
    run_state = _poll(
        "the reconciliation run",
        lambda: (
            (state := client.get(f"/api/v1/runs/{run_id}").json())["status"]
            in {"succeeded", "failed"}
            and state
        ),
        300,
    )
    if run_state["status"] != "succeeded":
        raise RuntimeError(
            f"run failed: {run_state.get('error_code')} {run_state.get('error_message')}"
        )
    _poll("post-run jobs to drain", lambda: stack.queue_busy() == 0, 300)
    cases = _all(client, "/api/v1/cases")

    export = client.post(
        "/api/v1/exports",
        json={"run_id": run_id, "filter": {"from": "2026-09-01", "to": "2026-09-30"}},
        headers={"Idempotency-Key": f"demo-export-{stack.tag}"},
    )
    _expect(export, 202)
    export_id = export.json()["id"]
    _step("exporting workbook")
    export_state = _poll(
        "the export",
        lambda: (
            (state := client.get(f"/api/v1/exports/{export_id}").json())["status"]
            in {"ready", "failed"}
            and state
        ),
        300,
    )
    workbook = b""
    if export_state["status"] == "ready":
        download = client.get(f"/api/v1/exports/{export_id}/download")
        _expect(download, 200)
        workbook = download.content
    # Home KPI cards for the same period as the export (they must equal its Summary sheet).
    kpis = client.get("/api/v1/kpis", params={"from": "2026-09-01", "to": "2026-09-30"})
    _expect(kpis, 200)
    return {
        "kpis": kpis.json(),
        "summary": summary,
        "files": files,
        "documents": documents,
        "run": run_state,
        "cases": cases,
        "export": export_state,
        "workbook": workbook,
    }


# ----------------------------------------------------------------- comparison


def compare(result: dict[str, Any]) -> list[str]:
    expected = yaml.safe_load((FIXTURE / "expected.yaml").read_text(encoding="utf-8"))
    problems: list[str] = []

    for item in result["files"]:
        # "partial" is a parsed file holding documents that need review (the unreadable scan).
        if item.get("status") not in {"parsed", "partial"}:
            problems.append(
                f"file {item.get('display_name')}: {item.get('status')} "
                f"{item.get('error_code') or ''}"
            )
    counts = Counter(str(document.get("doc_type")) for document in result["documents"])
    for doc_type, count in expected["documents"].items():
        if counts.get(doc_type, 0) != count:
            problems.append(
                f"documents[{doc_type}]: expected {count}, got {counts.get(doc_type, 0)}"
            )

    actual = [
        {
            "invoice": case.get("invoice_no"),
            "status": case.get("computed_status"),
            "reason_codes": sorted(case.get("reason_codes") or []),
            "po": (case.get("po_nos") or [None])[0],
            "at_risk": _decimal_text(case.get("at_risk")),
        }
        for case in result["cases"]
    ]
    remaining = list(actual)
    for want in expected["cases"]:
        match = next(
            (
                got
                for got in remaining
                if got["invoice"] == want["invoice"] and got["status"] == want["status"]
            ),
            None,
        )
        label = f"{want['scenario']} {want['invoice']}"
        if match is None:
            seen = [got for got in remaining if got["invoice"] == want["invoice"]]
            problems.append(f"{label}: expected {want['status']}, got {_brief(seen) or 'no case'}")
            continue
        remaining.remove(match)
        if match["reason_codes"] != sorted(want["reason_codes"]):
            problems.append(
                f"{label}: reasons {match['reason_codes']} != {sorted(want['reason_codes'])}"
            )
        if "po" in want and match["po"] != want["po"]:
            problems.append(f"{label}: linked to {match['po']}, expected {want['po']}")
        if "at_risk" in want and match["at_risk"] != want["at_risk"]:
            problems.append(f"{label}: at risk {match['at_risk']}, expected {want['at_risk']}")
    for extra in remaining:
        problems.append(f"unexpected case: {_brief([extra])}")

    total_at_risk = sum((Decimal(item["at_risk"] or "0") for item in actual), Decimal("0"))
    if f"{total_at_risk:.2f}" != expected["totals"]["at_risk"]:
        problems.append(f"total at risk {total_at_risk:.2f} != {expected['totals']['at_risk']}")

    if result["export"]["status"] != "ready":
        problems.append(
            f"export {result['export']['status']}: {result['export'].get('error_code')}"
        )
    else:
        problems += _check_workbook(
            result["workbook"], len(result["cases"]), total_at_risk, result["kpis"]
        )
    return problems


def _check_workbook(
    content: bytes, case_count: int, at_risk: Decimal, kpis: dict[str, Any]
) -> list[str]:
    from openpyxl import load_workbook

    problems: list[str] = []
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        expected_sheets = [
            "Summary",
            "Invoice Recon",
            "Line Recon",
            "Exceptions",
            "Vendor Summary",
        ]
        if workbook.sheetnames != expected_sheets:
            problems.append(f"workbook sheets {workbook.sheetnames}")
            return problems
        invoice_rows = [
            row
            for row in workbook["Invoice Recon"].iter_rows(min_row=2, values_only=True)
            if any(row)
        ]
        if len(invoice_rows) != case_count:
            problems.append(f"Invoice Recon rows {len(invoice_rows)} != cases {case_count}")
        summary = {
            str(row[0]): row[1]
            for row in workbook["Summary"].iter_rows(min_row=2, max_col=2, values_only=True)
            if row and row[0] is not None
        }
    finally:
        workbook.close()
    if summary.get("Invoices") != case_count:
        problems.append(f"Summary Invoices {summary.get('Invoices')} != cases {case_count}")
    if _decimal_text(summary.get("₹ At-risk")) != f"{at_risk:.2f}":
        problems.append(f"Summary At-risk {summary.get('₹ At-risk')} != cases {at_risk:.2f}")
    problems += _check_home_kpis(summary, kpis)
    return problems


def _check_home_kpis(summary: dict[str, Any], kpis: dict[str, Any]) -> list[str]:
    """Home KPI cards and the workbook Summary come from the same rule (App Flow §Home)."""
    invoices = summary.get("Invoices") or 0
    matched_pct = (
        None
        if not invoices
        else (Decimal(100) * Decimal(summary.get("Matched") or 0) / Decimal(invoices)).quantize(
            Decimal("0.01")
        )
    )
    pairs = [
        ("Invoices", invoices, kpis.get("invoice_count")),
        (
            "Matched %",
            _decimal_text(matched_pct),
            _decimal_text(kpis.get("matched_pct")),
        ),
        ("Pending", summary.get("Pending"), kpis.get("pending_count")),
        ("Exceptions", summary.get("Exceptions"), kpis.get("exception_count")),
        ("Unlinked", summary.get("Unlinked"), kpis.get("unlinked_count")),
        (
            "At-risk",
            _decimal_text(summary.get("₹ At-risk")),
            _decimal_text(kpis.get("at_risk_total")),
        ),
    ]
    return [
        f"Home KPI {name} {home} != workbook Summary {sheet}"
        for name, sheet, home in pairs
        if sheet != home
    ]


def _decimal_text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return f"{Decimal(str(value)):.2f}"


def _brief(cases: list[dict[str, Any]]) -> str:
    return "; ".join(
        f"{case['status']} {case['reason_codes']} po={case['po']} risk={case['at_risk']}"
        for case in cases
    )


# ------------------------------------------------------------------- commands


def command_render(out: Path) -> int:
    uploads = render_month(load_month(), out)
    for path, zone in uploads:
        print(f"{zone:8} {path.name} ({path.stat().st_size} bytes)")
    return 0


def command_run(*, build: bool, keep: bool, report: Path | None) -> int:
    month = load_month()
    if build:
        _step(f"building {DEMO_IMAGE} from the working tree")
        _docker(
            "build",
            "-f",
            str(ROOT / "deploy" / "backend.Dockerfile"),
            "-t",
            DEMO_IMAGE,
            str(ROOT),
        )
    if not _docker("images", "-q", MINIO_IMAGE, check=False):
        _step(f"building {MINIO_IMAGE}")
        _docker(
            "build",
            "-f",
            str(ROOT / "deploy" / "minio.Dockerfile"),
            "-t",
            MINIO_IMAGE,
            str(ROOT),
        )
    workdir = Path(tempfile.mkdtemp(prefix="reconbench-demo-"))
    stack = Stack(DEMO_IMAGE)
    started = time.monotonic()
    try:
        uploads = render_month(month, workdir / "upload")
        _step(f"rendered {len(uploads)} files; starting disposable stack {stack.tag}")
        stack.start()
        client = Client(f"http://127.0.0.1:{stack.api_port}")
        client.wait_ready()
        client.login(stack.email, stack.password)
        result = drive(client, stack, uploads, month)
        problems = compare(result)
        if report is not None:
            report.write_text(
                json.dumps(
                    {key: value for key, value in result.items() if key != "workbook"},
                    indent=2,
                    default=str,
                ),
                encoding="utf-8",
            )
        _print_cases(result["cases"])
        elapsed = time.monotonic() - started
        if problems:
            print(f"\nDEMO FAILED ({len(problems)} problems, {elapsed:.0f}s):")
            for problem in problems:
                print(f"  - {problem}")
            return 1
        print(f"\nDEMO PASSED: {len(result['cases'])} cases match expected.yaml ({elapsed:.0f}s)")
        return 0
    except Exception:
        for name in ("api", "worker"):
            if f"{stack.tag}-{name}" in stack.containers:
                print(f"\n--- {name} log tail ---\n{stack.logs(name)}", file=sys.stderr)
        raise
    finally:
        if keep:
            _step(f"kept stack {stack.tag} at http://127.0.0.1:{stack.api_port}")
        else:
            stack.stop()
        shutil.rmtree(workdir, ignore_errors=True)


def _print_cases(cases: list[dict[str, Any]]) -> None:
    print(
        "\ninvoice              status          reasons                         po          at risk"
    )
    for case in sorted(cases, key=lambda item: str(item.get("invoice_no"))):
        print(
            f"{case.get('invoice_no')!s:20} {case.get('computed_status')!s:15} "
            f"{','.join(case.get('reason_codes') or [])!s:31} "
            f"{(case.get('po_nos') or ['-'])[0]!s:11} {case.get('at_risk')}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    commands = parser.add_subparsers(dest="command", required=True)
    render = commands.add_parser("render", help="write the month's upload files")
    render.add_argument("--out", type=Path, required=True)
    run = commands.add_parser("run", help="full-stack smoke against a disposable stack")
    run.add_argument("--no-build", action="store_true", help=f"reuse the existing {DEMO_IMAGE}")
    run.add_argument("--keep", action="store_true", help="leave the stack running afterwards")
    run.add_argument("--report", type=Path, help="write API results (no workbook) as JSON")
    arguments = parser.parse_args()
    if arguments.command == "render":
        return command_render(arguments.out)
    return command_run(build=not arguments.no_build, keep=arguments.keep, report=arguments.report)


if __name__ == "__main__":
    raise SystemExit(main())
